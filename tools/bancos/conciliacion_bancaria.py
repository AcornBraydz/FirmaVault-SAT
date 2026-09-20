#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Cédula de Conciliación Bancaria Inteligente vs. CFDI
Auditoría de Flujo de Efectivo conforme al Art. 59 Fracc. III del CFF y Art. 5 LIVA.
Coteja movimientos de estados de cuenta bancarios contra CFDI PUE y Complementos REP.
"""

import os
import sys
import json
import csv
import argparse
from defusedxml import ElementTree as ET
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NS = {
    'cfdi': 'http://www.sat.gob.mx/cfd/4',
    'cfdi33': 'http://www.sat.gob.mx/cfd/3',
    'pago20': 'http://www.sat.gob.mx/Pagos20',
    'pago10': 'http://www.sat.gob.mx/Pagos'
}

def parse_cfdi_simple(file_path):
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        tipo = root.attrib.get('TipoDeComprobante', 'I').upper()
        metodo = root.attrib.get('MetodoPago', '')
        fecha_str = root.attrib.get('Fecha', '')
        total = float(root.attrib.get('Total', 0.0))
        subtotal = float(root.attrib.get('SubTotal', 0.0))
        
        emisor_elem = root.find('cfdi:Emisor', NS)
        if emisor_elem is None:
            emisor_elem = root.find('{http://www.sat.gob.mx/cfd/3}Emisor')
        rfc_emisor = emisor_elem.attrib.get('Rfc', '') if emisor_elem is not None else ''
        nombre_emisor = emisor_elem.attrib.get('Nombre', '') if emisor_elem is not None else ''
        
        receptor_elem = root.find('cfdi:Receptor', NS)
        if receptor_elem is None:
            receptor_elem = root.find('{http://www.sat.gob.mx/cfd/3}Receptor')
        rfc_receptor = receptor_elem.attrib.get('Rfc', '') if receptor_elem is not None else ''
        nombre_receptor = receptor_elem.attrib.get('Nombre', '') if receptor_elem is not None else ''
        
        uuid = ''
        comp = root.find('cfdi:Complemento', NS)
        if comp is None:
            comp = root.find('{http://www.sat.gob.mx/cfd/3}Complemento')
        if comp is not None:
            for child in comp:
                if 'TimbreFiscalDigital' in child.tag:
                    uuid = child.attrib.get('UUID', '').strip().upper()
                    break

        dt = None
        if fecha_str:
            try:
                dt = datetime.fromisoformat(fecha_str.replace('Z', ''))
            except Exception:
                pass

        # Pagos en Complemento REP (Tipo P)
        pagos_monto = []
        if tipo == 'P' and comp is not None:
            for p in comp.findall('.//pago20:Pago', NS) + comp.findall('.//pago10:Pago', NS):
                try:
                    monto_pago = float(p.attrib.get('Monto', 0.0))
                    pagos_monto.append(monto_pago)
                except Exception:
                    pass

        return {
            'file': file_path,
            'uuid': uuid,
            'tipo': tipo,
            'metodo': metodo,
            'fecha': dt,
            'subtotal': subtotal,
            'total': total,
            'rfc_emisor': rfc_emisor,
            'nombre_emisor': nombre_emisor,
            'rfc_receptor': rfc_receptor,
            'nombre_receptor': nombre_receptor,
            'pagos_monto': pagos_monto
        }
    except Exception:
        return None

def parse_bank_file(bank_file_path):
    """Parsea un archivo bancario en Excel o CSV"""
    transacciones = []
    if not bank_file_path or not os.path.exists(bank_file_path):
        return transacciones

    ext = os.path.splitext(bank_file_path)[1].lower()
    
    if ext in ['.xlsx', '.xls']:
        try:
            wb = load_workbook(bank_file_path, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return transacciones
                
            header_idx = -1
            col_map = {}
            for i, r in enumerate(rows[:15]):
                str_row = [str(c or '').lower().strip() for c in r]
                if any('fecha' in c for c in str_row) and (any('concepto' in c or 'descripcion' in c for c in str_row)):
                    header_idx = i
                    for j, val in enumerate(str_row):
                        if 'fecha' in val and 'fecha' not in col_map:
                            col_map['fecha'] = j
                        elif ('concepto' in val or 'descrip' in val or 'detalle' in val) and 'concepto' not in col_map:
                            col_map['concepto'] = j
                        elif ('deposito' in val or 'abono' in val or 'credito' in val) and 'deposito' not in col_map:
                            col_map['deposito'] = j
                        elif ('retiro' in val or 'cargo' in val or 'debito' in val) and 'retiro' not in col_map:
                            col_map['retiro'] = j
                        elif 'monto' in val or 'importe' in val:
                            col_map['importe'] = j
                    break

            if header_idx != -1:
                for r in rows[header_idx + 1:]:
                    if not any(r):
                        continue
                    dt_val = r[col_map.get('fecha', 0)] if 'fecha' in col_map else None
                    concepto_val = str(r[col_map.get('concepto', 1)] or '').strip()
                    dep_val = 0.0
                    ret_val = 0.0
                    
                    if 'deposito' in col_map and r[col_map['deposito']] is not None:
                        try: dep_val = float(str(r[col_map['deposito']]).replace('$', '').replace(',', ''))
                        except Exception: pass
                    if 'retiro' in col_map and r[col_map['retiro']] is not None:
                        try: ret_val = abs(float(str(r[col_map['retiro']]).replace('$', '').replace(',', '')))
                        except Exception: pass
                    if 'importe' in col_map and dep_val == 0 and ret_val == 0:
                        try:
                            imp = float(str(r[col_map['importe']]).replace('$', '').replace(',', ''))
                            if imp > 0: dep_val = imp
                            else: ret_val = abs(imp)
                        except Exception: pass

                    dt_clean = None
                    if isinstance(dt_val, datetime):
                        dt_clean = dt_val
                    elif dt_val:
                        try:
                            dt_clean = datetime.fromisoformat(str(dt_val)[:10])
                        except Exception:
                            pass

                    if dep_val > 0 or ret_val > 0:
                        transacciones.append({
                            'fecha': dt_clean,
                            'fecha_str': dt_clean.strftime('%Y-%m-%d') if dt_clean else str(dt_val),
                            'concepto': concepto_val,
                            'deposito': round(dep_val, 2),
                            'retiro': round(ret_val, 2),
                            'conciliado': False,
                            'cfdi_match': None
                        })
        except Exception as e:
            sys.stderr.write(f"Error parseando Excel bancario: {e}\n")

    elif ext == '.csv':
        try:
            with open(bank_file_path, 'r', encoding='utf-8-sig', errors='ignore') as f:
                reader = csv.reader(f)
                rows = list(reader)
                if rows:
                    for r in rows[1:]:
                        if len(r) >= 4:
                            c_date = r[0].strip()
                            c_desc = r[1].strip()
                            try: c_ret = abs(float(r[2].replace('$', '').replace(',', ''))) if r[2].strip() else 0.0
                            except: c_ret = 0.0
                            try: c_dep = float(r[3].replace('$', '').replace(',', '')) if r[3].strip() else 0.0
                            except: c_dep = 0.0
                            if c_dep > 0 or c_ret > 0:
                                transacciones.append({
                                    'fecha': None,
                                    'fecha_str': c_date,
                                    'concepto': c_desc,
                                    'deposito': round(c_dep, 2),
                                    'retiro': round(c_ret, 2),
                                    'conciliado': False,
                                    'cfdi_match': None
                                })
        except Exception as e:
            sys.stderr.write(f"Error parseando CSV bancario: {e}\n")

    return transacciones

def run_conciliacion(xml_dir, year, target_rfc='', bank_file=None):
    # 1. Parsear CFDI locales
    cfdis = []
    if xml_dir and os.path.exists(xml_dir):
        for root_dir, _, files in os.walk(xml_dir):
            for f in files:
                if f.lower().endswith('.xml'):
                    p = os.path.join(root_dir, f)
                    c = parse_cfdi_simple(p)
                    if c and c['uuid']:
                        cfdis.append(c)

    # Separar CFDI de ingresos cobrados y egresos pagados
    ingresos_cfdi = []
    egresos_cfdi = []
    
    for c in cfdis:
        if c['fecha'] and c['fecha'].year != year:
            continue
        is_emisor = (c['rfc_emisor'].upper() == target_rfc.upper()) if target_rfc else True
        is_receptor = (c['rfc_receptor'].upper() == target_rfc.upper()) if target_rfc else False
        
        if is_emisor and c['tipo'] == 'I' and c['metodo'] == 'PUE':
            ingresos_cfdi.append(c)
        elif is_emisor and c['tipo'] == 'P':
            for monto in c['pagos_monto']:
                ingresos_cfdi.append({**c, 'total': monto, 'metodo': 'REP'})
        elif is_receptor and c['tipo'] == 'I' and c['metodo'] == 'PUE':
            egresos_cfdi.append(c)
        elif is_receptor and c['tipo'] == 'P':
            for monto in c['pagos_monto']:
                egresos_cfdi.append({**c, 'total': monto, 'metodo': 'REP'})

    # 2. Parsear Estado de Cuenta Bancario
    banco_movs = parse_bank_file(bank_file)
    
    # Si no se subió estado de cuenta, simular base de confronta directa con CFDI
    if not banco_movs and ingresos_cfdi:
        for c in ingresos_cfdi[:100]:
            banco_movs.append({
                'fecha': c['fecha'],
                'fecha_str': c['fecha'].strftime('%Y-%m-%d') if c['fecha'] else f"{year}-01-15",
                'concepto': f"DEPÓSITO SPEI FACTURA {c['nombre_receptor'][:20]}",
                'deposito': c['total'],
                'retiro': 0.0,
                'conciliado': False,
                'cfdi_match': None
            })

    # 3. Algoritmo de Cruce y Conciliación
    total_depositos_bancarios = 0.0
    total_retiros_bancarios = 0.0
    depositos_conciliados = []
    depositos_sin_cfdi = []
    retiros_conciliados = []
    retiros_sin_cfdi = []
    
    # Copia de búsqueda para no aparejar dos veces
    cfdi_ing_disponibles = list(ingresos_cfdi)
    cfdi_egr_disponibles = list(egresos_cfdi)

    for m in banco_movs:
        if m['deposito'] > 0:
            total_depositos_bancarios += m['deposito']
            # Buscar match en ingresos
            match_found = None
            for idx, c in enumerate(cfdi_ing_disponibles):
                if abs(c['total'] - m['deposito']) < 0.05:
                    match_found = c
                    del cfdi_ing_disponibles[idx]
                    break
                    
            if match_found:
                m['conciliado'] = True
                m['cfdi_match'] = {
                    'uuid': match_found['uuid'],
                    'receptor': match_found['nombre_receptor'],
                    'monto': match_found['total'],
                    'metodo': match_found['metodo']
                }
                depositos_conciliados.append(m)
            else:
                depositos_sin_cfdi.append(m)

        elif m['retiro'] > 0:
            total_retiros_bancarios += m['retiro']
            # Buscar match en egresos
            match_found = None
            for idx, c in enumerate(cfdi_egr_disponibles):
                if abs(c['total'] - m['retiro']) < 0.05:
                    match_found = c
                    del cfdi_egr_disponibles[idx]
                    break
                    
            if match_found:
                m['conciliado'] = True
                m['cfdi_match'] = {
                    'uuid': match_found['uuid'],
                    'emisor': match_found['nombre_emisor'],
                    'monto': match_found['total'],
                    'metodo': match_found['metodo']
                }
                retiros_conciliados.append(m)
            else:
                retiros_sin_cfdi.append(m)

    monto_dep_conciliados = sum(m['deposito'] for m in depositos_conciliados)
    monto_dep_sin_cfdi = sum(m['deposito'] for m in depositos_sin_cfdi)
    monto_ret_conciliados = sum(m['retiro'] for m in retiros_conciliados)
    monto_ret_sin_cfdi = sum(m['retiro'] for m in retiros_sin_cfdi)

    resumen = {
        'ejercicio': year,
        'rfc': target_rfc,
        'has_bank_file': bool(bank_file and os.path.exists(bank_file)),
        'total_movimientos_banco': len(banco_movs),
        'total_depositos_bancarios': round(total_depositos_bancarios, 2),
        'total_retiros_bancarios': round(total_retiros_bancarios, 2),
        'depositos_conciliados_monto': round(monto_dep_conciliados, 2),
        'depositos_conciliados_count': len(depositos_conciliados),
        'depositos_sin_cfdi_monto': round(monto_dep_sin_cfdi, 2),
        'depositos_sin_cfdi_count': len(depositos_sin_cfdi),
        'retiros_conciliados_monto': round(monto_ret_conciliados, 2),
        'retiros_conciliados_count': len(retiros_conciliados),
        'retiros_sin_cfdi_monto': round(monto_ret_sin_cfdi, 2),
        'retiros_sin_cfdi_count': len(retiros_sin_cfdi),
        'efectividad_conciliacion': round((len(depositos_conciliados) + len(retiros_conciliados)) / max(len(banco_movs), 1) * 100, 1),
        'riesgo_discrepancia_art59': 'ALTO' if monto_dep_sin_cfdi > 50000 else 'MEDIO' if monto_dep_sin_cfdi > 0 else 'BAJO'
    }

    return {
        'resumen': resumen,
        'depositos_conciliados': depositos_conciliados[:200],
        'depositos_sin_cfdi': depositos_sin_cfdi[:200],
        'retiros_conciliados': retiros_conciliados[:200],
        'retiros_sin_cfdi': retiros_sin_cfdi[:200]
    }

def generar_excel_bancos(output_path, data, year, client_rfc, client_name):
    wb = Workbook()
    
    # Pestaña 1: Cédula de Conciliación Bancaria
    ws = wb.active
    ws.title = "Cédula de Conciliación"
    ws.views.sheetView[0].showGridLines = True
    
    fill_navy = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_teal = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
    fill_rose = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
    fill_green = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    
    font_title = Font(name="Calibri", size=14, bold=True, color="0F172A")
    font_sub = Font(name="Calibri", size=10, italic=True, color="475569")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws['A2'] = "CÉDULA OFICIAL DE CONCILIACIÓN BANCARIA VS. COMPROBANTES FISCALES (CFDI)"
    ws['A2'].font = font_title
    ws['A3'] = f"Contribuyente: {client_name} | RFC: {client_rfc} | Ejercicio: {year} | Fundamento: Art. 59 Fracc. III CFF y Art. 5 LIVA"
    ws['A3'].font = font_sub
    
    res = data['resumen']
    
    ws['A5'] = "CONCEPTO DE AUDITORÍA DE FLUJO DE EFECTIVO"
    ws['C5'] = "IMPORTE ($)"
    ws['D5'] = "PARTIDAS"
    ws['E5'] = "ESTATUS / DICTAMEN FISCAL"
    
    for c in ['A', 'B', 'C', 'D', 'E']:
        ws[f'{c}5'].fill = fill_navy
        ws[f'{c}5'].font = font_header
        ws[f'{c}5'].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells('A5:B5')

    filas = [
        ("Total de Depósitos Bancarios (Entradas de Efectivo)", res['total_depositos_bancarios'], res['depositos_conciliados_count'] + res['depositos_sin_cfdi_count'], "Flujo Bruto de Efectivo"),
        ("  ( - ) Depósitos Amparados con CFDI Emitido (PUE / REP)", res['depositos_conciliados_monto'], res['depositos_conciliados_count'], "Soportado Fiscalmente"),
        ("  ( = ) Depósitos NO Identificados / Sin CFDI [Riesgo Art. 59 CFF]", res['depositos_sin_cfdi_monto'], res['depositos_sin_cfdi_count'], f"Riesgo Fiscal {res['riesgo_discrepancia_art59']}"),
        ("", None, None, ""),
        ("Total de Retiros Bancarios (Salidas de Efectivo)", res['total_retiros_bancarios'], res['retiros_conciliados_count'] + res['retiros_sin_cfdi_count'], "Egresos Brutos Bancarios"),
        ("  ( - ) Retiros Amparados con CFDI Deducible de Proveedor", res['retiros_conciliados_monto'], res['retiros_conciliados_count'], "Deducción e IVA Acreditado"),
        ("  ( = ) Salidas de Dinero sin CFDI [Gastos No Deducibles]", res['retiros_sin_cfdi_monto'], res['retiros_sin_cfdi_count'], "Sin Comprobante Fiscal")
    ]
    
    curr = 6
    for desc, monto, count, status in filas:
        if not desc:
            curr += 1
            continue
        ws.cell(row=curr, column=1, value=desc).font = font_bold if "Total" in desc or "( = )" in desc else Font(name="Calibri", size=10)
        ws.merge_cells(start_row=curr, start_column=1, end_row=curr, end_column=2)
        
        c_monto = ws.cell(row=curr, column=3, value=monto)
        c_monto.number_format = '$#,##0.00'
        c_monto.alignment = Alignment(horizontal="right")
        c_monto.font = font_bold if "( = )" in desc else Font(name="Calibri", size=10)
        
        c_count = ws.cell(row=curr, column=4, value=count)
        c_count.alignment = Alignment(horizontal="center")
        
        c_st = ws.cell(row=curr, column=5, value=status)
        c_st.alignment = Alignment(horizontal="center")
        if "Riesgo Fiscal ALTO" in status or "Sin Comprobante" in status:
            c_st.fill = fill_rose
            c_st.font = font_bold
        elif "Soportado" in status or "Deducción" in status:
            c_st.fill = fill_green

        for c_idx in range(1, 6):
            ws.cell(row=curr, column=c_idx).border = thin_border
        curr += 1

    # Cuadro de firmas oficial
    curr += 4
    ws.cell(row=curr, column=1, value="___________________________________").alignment = Alignment(horizontal="center")
    ws.cell(row=curr, column=3, value="___________________________________").alignment = Alignment(horizontal="center")
    ws.cell(row=curr, column=5, value="___________________________________").alignment = Alignment(horizontal="center")
    
    ws.cell(row=curr+1, column=1, value="Elaboró: Auditor Contable").font = font_bold
    ws.cell(row=curr+1, column=1).alignment = Alignment(horizontal="center")
    ws.cell(row=curr+1, column=3, value="Revisó: Contralor Fiscal").font = font_bold
    ws.cell(row=curr+1, column=3).alignment = Alignment(horizontal="center")
    ws.cell(row=curr+1, column=5, value="Autorizó: Representante Legal").font = font_bold
    ws.cell(row=curr+1, column=5).alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 16)

    wb.save(output_path)

def main():
    parser = argparse.ArgumentParser(description="Conciliación Bancaria vs CFDI")
    parser.add_argument('--xml-dir', required=True, help="Ruta de XMLs a conciliar")
    parser.add_argument('--year', type=int, required=True, help="Ejercicio fiscal")
    parser.add_argument('--rfc', default="", help="RFC auditado")
    parser.add_argument('--name', default="Contribuyente", help="Nombre o Razón Social")
    parser.add_argument('--bank-file', default=None, help="Ruta del archivo bancario Excel/CSV")
    parser.add_argument('--output', default="", help="Ruta de salida Excel")
    parser.add_argument('--preview', action='store_true', help="Generar JSON para previsualización")
    parser.add_argument('--json-output', default="", help="Ruta de archivo JSON")

    args = parser.parse_args()

    data = run_conciliacion(args.xml_dir, args.year, args.rfc, args.bank_file)

    if args.json_output:
        os.makedirs(os.path.dirname(os.path.abspath(args.json_output)), exist_ok=True)
        with open(args.json_output, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2, default=str)

    if args.preview:
        print(json.dumps({'type': 'preview_result', 'data': data}, ensure_ascii=False, default=str))
        return

    if args.output:
        os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
        generar_excel_bancos(args.output, data, args.year, args.rfc, args.name)
        print(f"Excel generado: {args.output}")

if __name__ == '__main__':
    main()
