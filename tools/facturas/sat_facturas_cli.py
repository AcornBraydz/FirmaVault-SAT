import sys
import os

# Configurar UTF-8 para stdout y stderr en Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace', line_buffering=True)
        sys.stderr.reconfigure(encoding='utf-8', errors='replace', line_buffering=True)
    except Exception:
        pass

import json
import re
import time
import argparse
import tempfile
import calendar
import shutil
import zipfile
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import subprocess

UUID_REGEX = re.compile(r'^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$')
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.action_chains import ActionChains

SAT_LOGIN_URL = "https://portalcfdi.facturaelectronica.sat.gob.mx/"
URL_EMISOR = "https://portalcfdi.facturaelectronica.sat.gob.mx/ConsultaEmisor.aspx"
URL_RECEPTOR = "https://portalcfdi.facturaelectronica.sat.gob.mx/ConsultaReceptor.aspx"

MESES_NOMBRES = [
    "", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"
]

_last_emitted_progress = 0

def emit(data):
    """Emite evento estructurado en formato JSON a stdout garantizando no retroceso en progreso"""
    global _last_emitted_progress
    try:
        if isinstance(data, dict) and data.get("type") == "progress":
            val = int(data.get("value", 0))
            if val < _last_emitted_progress:
                val = _last_emitted_progress
            else:
                _last_emitted_progress = val
            data["value"] = val

        print(json.dumps(data, ensure_ascii=False), flush=True)
    except Exception:
        print(json.dumps(data, ensure_ascii=True), flush=True)

def find_browser_executable():
    """Detecta la ruta del navegador disponible en el sistema (Brave, Chrome o Edge)"""
    candidate_paths = [
        r"C:\Program Files\BraveSoftware\Brave-Browser\Application\brave.exe",
        os.path.expandvars(r"%LOCALAPPDATA%\BraveSoftware\Brave-Browser\Application\brave.exe"),
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        os.path.expandvars(r"%LOCALAPPDATA%\Google\Chrome\Application\chrome.exe"),
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    ]
    for p in candidate_paths:
        if os.path.exists(p):
            return p
    return None

def format_fecha_display(fecha_raw):
    """Formatea la fecha de emisión a formato legible (DD-Mes-AAAA)"""
    if not fecha_raw:
        return ""
    try:
        # Ejemplo: 2024-05-15T10:30:00 o 2024-05-15
        clean_date = fecha_raw.split("T")[0]
        dt = datetime.strptime(clean_date, "%Y-%m-%d")
        mes_str = MESES_NOMBRES[dt.month]
        return f"{dt.day:02d}-{mes_str}-{dt.year}"
    except Exception:
        return fecha_raw

def parse_monto(val_str):
    """Limpia y convierte monto a float"""
    if not val_str:
        return 0.0
    try:
        cleaned = str(val_str).replace("$", "").replace(",", "").strip()
        return round(float(cleaned), 2)
    except Exception:
        return 0.0

def poblar_hoja_excel(ws, registros_hoja, is_emitidas, header_color):
    """Llena una hoja de Excel con formato corporativo top-tier, fórmulas e importes contables"""
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    headers = [
        "Mes",
        "Folio",
        "Fecha de Emisión",
        "Cliente / Receptor" if is_emitidas else "Proveedor / Emisor",
        "RFC",
        "Importe",
        "IVA al 16%",
        "Total",
        "Estado del Comprobante",
        "Efecto del Comprobante"
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    fill_vigente = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    font_vigente = Font(name="Calibri", size=10, bold=True, color="047857")
    fill_cancelado = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    font_cancelado = Font(name="Calibri", size=10, bold=True, color="B91C1C")

    # Llenar datos
    for row_idx, r in enumerate(registros_hoja, 2):
        ws.cell(row=row_idx, column=1, value=r.get("Mes", "")).alignment = align_center
        ws.cell(row=row_idx, column=2, value=r.get("Folio", "")).alignment = align_center
        ws.cell(row=row_idx, column=3, value=r.get("Fecha de Emisión", "")).alignment = align_center
        ws.cell(row=row_idx, column=4, value=r.get("Nombre o Razón Social", "")).alignment = align_left
        ws.cell(row=row_idx, column=5, value=r.get("RFC", "")).alignment = align_center

        efecto_norm = clasificar_efecto(r.get("Efecto del Comprobante", ""))
        es_cero_iva = efecto_norm in ["nomina", "pago", "traslado"]
        real_subtotal = r.get("RealSubtotal")
        real_iva = r.get("RealIva")

        if real_subtotal is not None and real_iva is not None:
            cell_importe = ws.cell(row=row_idx, column=6, value=float(real_subtotal))
            cell_iva = ws.cell(row=row_idx, column=7, value=float(real_iva))
        elif es_cero_iva:
            cell_importe = ws.cell(row=row_idx, column=6, value=float(r.get("Total", 0.0)))
            cell_iva = ws.cell(row=row_idx, column=7, value=0.0)
        else:
            # Celda F (col 6): Fórmula Importe = ROUND(H{row}/1.16, 2)
            cell_importe = ws.cell(row=row_idx, column=6, value=f"=ROUND(H{row_idx}/1.16, 2)")
            # Celda G (col 7): Fórmula IVA = ROUND(F{row}*0.16, 2)
            cell_iva = ws.cell(row=row_idx, column=7, value=f"=ROUND(F{row_idx}*0.16, 2)")

        cell_importe.alignment = align_right
        cell_importe.number_format = '"$"#,##0.00'
        cell_iva.alignment = align_right
        cell_iva.number_format = '"$"#,##0.00'

        # Celda H (col 8): Total numérico
        total_val = float(r.get("Total", 0.0))
        cell_total = ws.cell(row=row_idx, column=8, value=total_val)
        cell_total.alignment = align_right
        cell_total.number_format = '"$"#,##0.00'

        estado = str(r.get("Estado del Comprobante", ""))
        cell_estado = ws.cell(row=row_idx, column=9, value=estado)
        cell_estado.alignment = align_center
        if "vigente" in estado.lower():
            cell_estado.fill = fill_vigente
            cell_estado.font = font_vigente
        elif "cancelado" in estado.lower():
            cell_estado.fill = fill_cancelado
            cell_estado.font = font_cancelado

        ws.cell(row=row_idx, column=10, value=r.get("Efecto del Comprobante", "")).alignment = align_center

        for col_i in range(1, 11):
            ws.cell(row=row_idx, column=col_i).border = thin_border
        ws.row_dimensions[row_idx].height = 20

    # Fila de Totales al final
    total_row_idx = len(registros_hoja) + 2
    fill_totales = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    font_totales = Font(name="Calibri", size=11, bold=True, color="0F172A")
    thick_top_double_bottom = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="94A3B8"),
        bottom=Side(style="double", color="0F172A")
    )

    for c in range(1, 6):
        cell = ws.cell(row=total_row_idx, column=c)
        cell.fill = fill_totales
        cell.border = thick_top_double_bottom
    label_total = ws.cell(row=total_row_idx, column=5, value="TOTAL:")
    label_total.font = font_totales
    label_total.alignment = align_right

    if len(registros_hoja) > 0:
        cell_tot_imp = ws.cell(row=total_row_idx, column=6, value=f"=SUM(F2:F{total_row_idx-1})")
        cell_tot_iva = ws.cell(row=total_row_idx, column=7, value=f"=SUM(G2:G{total_row_idx-1})")
        cell_tot_tot = ws.cell(row=total_row_idx, column=8, value=f"=SUM(H2:H{total_row_idx-1})")
    else:
        cell_tot_imp = ws.cell(row=total_row_idx, column=6, value=0.0)
        cell_tot_iva = ws.cell(row=total_row_idx, column=7, value=0.0)
        cell_tot_tot = ws.cell(row=total_row_idx, column=8, value=0.0)

    for cell in [cell_tot_imp, cell_tot_iva, cell_tot_tot]:
        cell.font = font_totales
        cell.alignment = align_right
        cell.number_format = '"$"#,##0.00'
        cell.fill = fill_totales
        cell.border = thick_top_double_bottom

    cell_conteo = ws.cell(row=total_row_idx, column=9, value=f"{len(registros_hoja)} comprobantes")
    cell_conteo.font = font_totales
    cell_conteo.alignment = align_center
    cell_conteo.fill = fill_totales
    cell_conteo.border = thick_top_double_bottom

    cell_empty10 = ws.cell(row=total_row_idx, column=10, value="")
    cell_empty10.fill = fill_totales
    cell_empty10.border = thick_top_double_bottom
    ws.row_dimensions[total_row_idx].height = 24

    # Auto-ajustar ancho de columnas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for c in col:
            val_str = str(c.value or "")
            if val_str.startswith("="):
                val_str = "$1,234,567.89"
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)


def clasificar_efecto(efecto_raw):
    """
    Normaliza el efecto del comprobante a una clave estándar:
    'ingreso', 'pago', 'egreso', 'nomina', 'traslado', 'otro'
    """
    txt = str(efecto_raw or "").lower().strip()
    if any(k in txt for k in ["nomina", "nómina"]) or txt == "n":
        return "nomina"
    if any(k in txt for k in ["pago"]) or txt == "p":
        return "pago"
    if any(k in txt for k in ["egreso", "devoluci", "crédito", "credito"]) or txt == "e":
        return "egreso"
    if any(k in txt for k in ["traslado"]) or txt == "t":
        return "traslado"
    if any(k in txt for k in ["ingreso"]) or txt == "i":
        return "ingreso"
    return "otro"


def generar_excel_facturas(registros, tipo, output_excel_path):
    """
    Genera el archivo Excel estructurado dinámicamente:
    - Hoja 1: Concentrado General ("Todas las Emitidas" o "Todas las Recibidas")
    - Hojas dinámicas creadas únicamente si existen comprobantes de ese tipo:
      * "Ingresos": Facturas de ingreso / comerciales
      * "Pagos (REP)": Complementos de recepción de pagos (ej. Telmex)
      * "Notas de Crédito": Comprobantes de egreso / devoluciones
      * "Nóminas": Recibos de nómina timbrados
      * "Traslados": Cartas porte y traslados
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_excel_path)), exist_ok=True)
    wb = openpyxl.Workbook()

    is_emitidas = tipo == "emitidas"

    # HOJA 1: Concentrado General (Siempre presente con el 100% de los comprobantes)
    ws_todas = wb.active
    ws_todas.title = "Todas las Emitidas" if is_emitidas else "Todas las Recibidas"
    color_todas = "312E81" if is_emitidas else "134E4A"
    poblar_hoja_excel(ws_todas, registros, is_emitidas, color_todas)

    # Definición de pestañas por Efecto fiscal
    efectos_config = [
        ("ingreso", "Ingresos", "065F46"),
        ("pago", "Pagos (REP)", "0F766E"),
        ("egreso", "Notas de Crédito", "B45309"),
        ("nomina", "Nóminas", "1E3A8A"),
        ("traslado", "Traslados", "475569"),
        ("otro", "Otros Comprobantes", "64748B"),
    ]

    for ef_key, sheet_title, color_header in efectos_config:
        sub_regs = [r for r in registros if clasificar_efecto(r.get("Efecto del Comprobante")) == ef_key]
        if sub_regs:
            ws_sub = wb.create_sheet(title=sheet_title)
            poblar_hoja_excel(ws_sub, sub_regs, is_emitidas, color_header)

    wb.save(output_excel_path)
    return output_excel_path

def extraer_modulo_facturas(driver, tipo, year, month, rfc, download_files, download_dir, emit, progress_range=(50, 90)):
    registros = []
    target_url = URL_EMISOR if tipo == "emitidas" else URL_RECEPTOR
    modulo_nombre = "Facturas Emitidas (Ingresos)" if tipo == "emitidas" else "Facturas Recibidas (Gastos)"
    emit({"type": "log", "msg": f"ðŸ§­ Navegando a {modulo_nombre}..."})

    try:
        driver.get(target_url)
    except Exception:
        pass
    time.sleep(1.5)

    wait = WebDriverWait(driver, 20)

    # Seleccionar búsqueda por fechas de forma segura con espera de postback
    try:
        wait.until(EC.presence_of_element_located((By.ID, "ctl00_MainContent_RdoFechas")))
        rdo_fechas = driver.find_element(By.ID, "ctl00_MainContent_RdoFechas")
        try:
            rdo_fechas.click()
        except Exception:
            driver.execute_script("arguments[0].click();", rdo_fechas)

        if tipo == "recibidas":
            try:
                WebDriverWait(driver, 15).until(
                    lambda d: bool(d.find_elements(By.ID, "DdlAnio")) and
                              not d.find_element(By.ID, "DdlAnio").get_attribute("disabled")
                )
            except Exception:
                driver.execute_script("__doPostBack('ctl00$MainContent$RdoFechas','');")
                WebDriverWait(driver, 15).until(
                    lambda d: bool(d.find_elements(By.ID, "DdlAnio")) and
                              not d.find_element(By.ID, "DdlAnio").get_attribute("disabled")
                )
        else:
            time.sleep(1.0)
            try:
                WebDriverWait(driver, 15).until(
                    lambda d: bool(d.find_elements(By.ID, "ctl00_MainContent_CldFechaInicial2_Calendario_text")) or
                              bool(d.find_elements(By.ID, "DdlAnio"))
                )
            except Exception:
                driver.execute_script("__doPostBack('ctl00$MainContent$RdoFechas','');")
                time.sleep(1.5)
    except Exception:
        pass

    is_all_months = str(month).lower() == "all"
    int_year = int(year)
    now_y = datetime.now().year
    now_m = datetime.now().month

    if is_all_months:
        # En el año en curso solo se auditan meses transcurridos; en años anteriores los 12 meses
        if int_year == now_y:
            target_months = list(range(1, now_m + 1))
        elif int_year > now_y:
            target_months = []
        else:
            target_months = list(range(1, 13))
    else:
        target_months = [int(month) if str(month).isdigit() else 1]

    total_months = len(target_months)
    if total_months == 0:
        emit({"type": "log", "msg": f"ℹ️ï¸ El ejercicio {int_year} no tiene meses fiscales transcurridos."})
        return registros, None, 0

    main_win = driver.current_window_handle
    comprobantes_dir = os.path.join(download_dir, f"Comprobantes_{tipo.upper()}_{rfc}_{int_year}")
    if download_files:
        os.makedirs(comprobantes_dir, exist_ok=True)

    for m_idx, int_month in enumerate(target_months):
        nombre_mes = MESES_NOMBRES[int_month]
        p_span = progress_range[1] - progress_range[0]
        p_val = int(progress_range[0] + (m_idx / total_months) * p_span)
        emit({"type": "progress", "value": p_val})
        emit({"type": "log", "msg": f"ðŸ” [{nombre_mes} {int_year}] Consultando comprobantes oficiales de {tipo}..."})

        # Asegurar radio de fechas activo
        try:
            rdo_fechas = driver.find_element(By.ID, "ctl00_MainContent_RdoFechas")
            if not rdo_fechas.is_selected():
                try:
                    rdo_fechas.click()
                except Exception:
                    driver.execute_script("arguments[0].click();", rdo_fechas)
                time.sleep(0.3)
        except Exception:
            pass

        if tipo == "recibidas":
            wait.until(EC.presence_of_element_located((By.ID, "DdlAnio")))
            ddl_a = driver.find_element(By.ID, "DdlAnio")
            driver.execute_script("arguments[0].removeAttribute('disabled');", ddl_a)
            Select(ddl_a).select_by_value(str(int_year))
            time.sleep(0.2)
            ddl_m = driver.find_element(By.ID, "ctl00_MainContent_CldFecha_DdlMes")
            driver.execute_script("arguments[0].removeAttribute('disabled');", ddl_m)
            Select(ddl_m).select_by_value(str(int_month))
            time.sleep(0.2)
        else:
            ddl_anio_list = driver.find_elements(By.ID, "DdlAnio")
            if ddl_anio_list and ddl_anio_list[0].is_displayed():
                driver.execute_script("arguments[0].removeAttribute('disabled');", ddl_anio_list[0])
                Select(ddl_anio_list[0]).select_by_value(str(int_year))
                time.sleep(0.2)
                ddl_mes_list = driver.find_elements(By.ID, "ctl00_MainContent_CldFecha_DdlMes")
                if ddl_mes_list and ddl_mes_list[0].is_displayed():
                    driver.execute_script("arguments[0].removeAttribute('disabled');", ddl_mes_list[0])
                    Select(ddl_mes_list[0]).select_by_value(str(int_month))
                    time.sleep(0.2)
            else:
                last_day = calendar.monthrange(int_year, int_month)[1]
                fecha_ini_str = f"01/{int_month:02d}/{int_year}"
                fecha_fin_str = f"{last_day:02d}/{int_month:02d}/{int_year}"

                driver.execute_script(f"""
                    var i1 = document.getElementById('ctl00_MainContent_CldFechaInicial2_Calendario_text');
                    var i2 = document.getElementById('ctl00_MainContent_CldFechaFinal2_Calendario_text');
                    if (i1) {{ i1.removeAttribute('disabled'); i1.value = '{fecha_ini_str}'; }}
                    if (i2) {{ i2.removeAttribute('disabled'); i2.value = '{fecha_fin_str}'; }}
                """)

                inputs_ini = driver.find_elements(By.XPATH, "//input[contains(@id, 'FechaInicial') or contains(@id, 'txtFechaIni') or contains(@id, 'Fecha_Ini')]")
                inputs_fin = driver.find_elements(By.XPATH, "//input[contains(@id, 'FechaFinal') or contains(@id, 'txtFechaFin') or contains(@id, 'Fecha_Fin')]")

                if inputs_ini and inputs_fin:
                    driver.execute_script(f"arguments[0].removeAttribute('disabled'); arguments[0].value = '{fecha_ini_str}';", inputs_ini[0])
                    driver.execute_script(f"arguments[0].removeAttribute('disabled'); arguments[0].value = '{fecha_fin_str}';", inputs_fin[0])
                    time.sleep(0.2)

        # Limpiar elementos previos en el DOM para que WebDriverWait no confunda resultados del mes previo
        try:
            driver.execute_script("""
                var oldTbl = document.getElementById('ctl00_MainContent_tblResult');
                if (oldTbl) oldTbl.id = 'ctl00_MainContent_tblResult_old';
                var oldMsg = document.getElementById('ctl00_MainContent_lblMensajeDisp');
                if (oldMsg) oldMsg.id = 'ctl00_MainContent_lblMensajeDisp_old';
            """)
        except Exception:
            pass

        # Ejecutar Búsqueda de forma robusta
        btn_buscar = None
        for _ in range(3):
            try:
                btn_buscar = WebDriverWait(driver, 5).until(
                    EC.presence_of_element_located((By.XPATH, "//input[@id='ctl00_MainContent_BtnBusqueda' or contains(@id, 'BtnBusqueda') or @value='Buscar CFDI']"))
                )
                if btn_buscar:
                    break
            except Exception:
                time.sleep(0.5)

        if not btn_buscar:
            btns = driver.find_elements(By.XPATH, "//input[contains(@id, 'BtnBusqueda') or @value='Buscar CFDI']")
            if btns:
                btn_buscar = btns[0]

        if btn_buscar:
            try:
                btn_buscar.click()
            except Exception:
                driver.execute_script("arguments[0].click();", btn_buscar)
        else:
            # Disparar directamente postback ASP.NET si el elemento se desasoció temporalmente del DOM
            try:
                driver.execute_script("""
                    if (typeof ocultaResultados === 'function') ocultaResultados();
                    if (typeof WebForm_DoPostBackWithOptions === 'function') {
                        WebForm_DoPostBackWithOptions(new WebForm_PostBackOptions("ctl00$MainContent$BtnBusqueda", "", true, "FolioFiscal", "", false, false));
                    } else if (typeof __doPostBack === 'function') {
                        __doPostBack('ctl00$MainContent$BtnBusqueda','');
                    }
                """)
            except Exception:
                pass

        # Esperar a que el servidor del SAT inyecte la NUEVA respuesta en el DOM
        try:
            WebDriverWait(driver, 10).until(
                lambda d: bool(d.find_elements(By.ID, "ctl00_MainContent_tblResult")) or
                          bool(d.find_elements(By.ID, "ctl00_MainContent_lblMensajeDisp"))
            )
            time.sleep(0.4)
        except Exception:
            time.sleep(1.0)

        # Revisar si hay mensaje de no encontrado
        try:
            lbl_mensaje = driver.find_elements(By.ID, "ctl00_MainContent_lblMensajeDisp")
            if lbl_mensaje and lbl_mensaje[0].is_displayed() and lbl_mensaje[0].text.strip() != "":
                emit({"type": "log", "msg": f"ℹ️ï¸ [{nombre_mes}] Sin facturas registradas en este periodo."})
                continue
        except Exception:
            pass

        # Extracción atómica y de ultra alta velocidad vía JavaScript (inmune a StaleElementReference)
        raw_rows = driver.execute_script("""
            var tbl = document.getElementById('ctl00_MainContent_tblResult');
            if (!tbl) return [];
            var trs = tbl.querySelectorAll('tbody tr');
            var out = [];
            for (var i = 0; i < trs.length; i++) {
                var tr = trs[i];
                if (tr.innerText.indexOf('RFC Emisor') !== -1) continue;
                var tds = tr.querySelectorAll('td');
                if (tds.length >= 11) {
                    var rData = [];
                    for (var j = 0; j < tds.length; j++) {
                        rData.push(tds[j].innerText.trim());
                    }
                    out.push(rData);
                }
            }
            return out;
        """)

        if not raw_rows:
            emit({"type": "log", "msg": f"ℹ️ï¸ [{nombre_mes}] 0 facturas encontradas."})
            continue

        mes_count = 0
        comprobantes_mes = []

        for celdas in raw_rows:
            # Localizar columna del Folio Fiscal (UUID) dinámicamente para soportar cualquier orden de iconos del SAT
            uuid_idx = -1
            for idx, c in enumerate(celdas):
                if UUID_REGEX.match(c):
                    uuid_idx = idx
                    break

            if uuid_idx == -1:
                continue

            folio = celdas[uuid_idx]
            rfc_emisor = celdas[uuid_idx + 1] if len(celdas) > uuid_idx + 1 else ""
            nombre_emisor = celdas[uuid_idx + 2] if len(celdas) > uuid_idx + 2 else ""
            rfc_receptor = celdas[uuid_idx + 3] if len(celdas) > uuid_idx + 3 else ""
            nombre_receptor = celdas[uuid_idx + 4] if len(celdas) > uuid_idx + 4 else ""
            fecha_emision_raw = celdas[uuid_idx + 5] if len(celdas) > uuid_idx + 5 else ""
            total_raw = celdas[uuid_idx + 8] if len(celdas) > uuid_idx + 8 else "0"
            efecto = celdas[uuid_idx + 9] if len(celdas) > uuid_idx + 9 else "Ingreso"
            estado = celdas[uuid_idx + 11] if len(celdas) > uuid_idx + 11 else "Vigente"

            total_num = parse_monto(total_raw)

            # Normalización fiscal rigurosa para detección de Nóminas y Servicios Exentos (Art. 15 LIVA)
            import unicodedata
            def _norm_str(s):
                return unicodedata.normalize('NFKD', str(s or '')).encode('ASCII', 'ignore').decode('utf-8').lower().strip()

            efecto_norm = _norm_str(efecto)
            nombre_emisor_norm = _norm_str(nombre_emisor)
            nombre_receptor_norm = _norm_str(nombre_receptor)
            nombre_check = nombre_emisor_norm if tipo == "recibidas" else nombre_receptor_norm

            is_nomina = 'nomina' in efecto_norm or efecto_norm == 'n' or 'nmina' in efecto_norm
            is_educacion = any(w in nombre_check for w in ['instituto', 'colegio', 'escuela', 'universidad', 'educat', 'educac', 'ensenan', 'liceo', 'kinder', 'primaria', 'secundaria', 'preparatoria', 'academ', 'formacion', 'docen', 'pedagog']) and not ('seguro social' in nombre_check or 'imss' in nombre_check or 'issste' in nombre_check)
            is_medico = any(w in nombre_check for w in ['hospital', 'clinica', 'medico', 'sanatorio', 'salud', 'laboratorio', 'diagnostico'])
            is_pago = efecto_norm == 'p' or 'pago' in efecto_norm
            is_egreso = efecto_norm == 'e' or 'egreso' in efecto_norm

            if is_nomina or is_educacion or is_medico or is_pago:
                importe_calc = total_num
                iva_calc = 0.0
            else:
                importe_calc = round(total_num / 1.16, 2)
                iva_calc = round(total_num - importe_calc, 2)

            if tipo == "emitidas":
                rfc_final = rfc_receptor
                nombre_final = nombre_receptor
            else:
                rfc_final = rfc_emisor
                nombre_final = nombre_emisor

            metodo_pago_calc = "REP" if is_pago else ("Egreso" if is_egreso else ("Nómina" if is_nomina else "PUE"))
            forma_pago_calc = "REP" if is_pago else ("EGRESO" if is_egreso else ("99" if is_nomina else "SIN_XML"))

            reg_item = {
                "Mes": nombre_mes,
                "NumMes": int_month,
                "Folio": folio,
                "Fecha de Emisión": format_fecha_display(fecha_emision_raw),
                "fecha_emision": format_fecha_display(fecha_emision_raw),
                "Nombre o Razón Social": nombre_final,
                "nombre_razon_social": nombre_final,
                "RFC": rfc_final,
                "Importe": importe_calc,
                "IVA al 16%": iva_calc,
                "Total": total_num,
                "Estado del Comprobante": estado,
                "Efecto del Comprobante": efecto,
                "MetodoPago": metodo_pago_calc,
                "FormaPago": forma_pago_calc
            }
            registros.append(reg_item)
            comprobantes_mes.append((folio, reg_item))
            mes_count += 1

        emit({"type": "log", "msg": f"ðŸ“„ [{nombre_mes}] {mes_count} facturas oficiales detectadas."})

        # Descarga física de XML y PDF solo si download_files es True
        if download_files and len(comprobantes_mes) > 0:
            mes_folder = f"{int_month:02d}_{nombre_mes}"
            mes_dir = os.path.join(comprobantes_dir, mes_folder)
            xml_dir = os.path.join(mes_dir, "XML")
            pdf_dir = os.path.join(mes_dir, "PDF")
            os.makedirs(xml_dir, exist_ok=True)
            os.makedirs(pdf_dir, exist_ok=True)

            total_en_mes = len(comprobantes_mes)
            month_slice = p_span / total_months
            for idx, (folio, item) in enumerate(comprobantes_mes):
                file_p = int(p_val + ((idx + 1) / total_en_mes) * (month_slice * 0.95))
                emit({"type": "progress", "value": file_p})
                if (idx + 1) % 10 == 0 or (idx + 1) == total_en_mes:
                    emit({"type": "log", "msg": f"ðŸ“¥ [{nombre_mes}] Descargando comprobantes oficiales: {idx + 1}/{total_en_mes} (.XML y .PDF)..."})
                xml_descargado = False
                pdf_descargado = False

                # 1. XML Oficial
                try:
                    xpath_xml = f"//table[@id='ctl00_MainContent_tblResult']//tr[contains(., '{folio}')]//span[contains(@id, 'BtnDescarga')] | //table[@id='ctl00_MainContent_tblResult']//tr[contains(., '{folio}')]//a[contains(@onclick, 'AccionCfdi')]"
                    btn_xml = driver.find_elements(By.XPATH, xpath_xml)
                    if btn_xml:
                        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", btn_xml[0])
                        time.sleep(0.04)
                        driver.execute_script("arguments[0].click();", btn_xml[0])
                        for _ in range(4):
                            time.sleep(0.2)
                            for fname in os.listdir(download_dir):
                                if fname.lower().endswith(".xml") and not fname.lower().endswith(".tmp"):
                                    src_f = os.path.join(download_dir, fname)
                                    dst_f = os.path.join(xml_dir, f"{folio}.xml")
                                    try:
                                        shutil.move(src_f, dst_f)
                                        xml_descargado = True
                                        try:
                                            from defusedxml import ElementTree as ET_sat
                                            xtree = ET_sat.parse(dst_f)
                                            xroot = xtree.getroot()
                                            x_sub = float(xroot.attrib.get('SubTotal', item.get('Total', 0.0)))
                                            x_desc = float(xroot.attrib.get('Descuento', 0.0))
                                            x_iva = 0.0
                                            for xel in xroot.iter():
                                                xtag = xel.tag.split('}')[-1] if '}' in xel.tag else xel.tag
                                                if xtag == 'Traslado':
                                                    ximp = xel.attrib.get('Impuesto', '')
                                                    if ximp in ['002', 'IVA']:
                                                        x_iva += float(xel.attrib.get('Importe', 0.0))
                                            item['RealSubtotal'] = round(x_sub - x_desc, 2)
                                            item['RealIva'] = round(x_iva, 2)
                                            x_forma = xroot.attrib.get('FormaPago', '')
                                            x_metodo = xroot.attrib.get('MetodoPago', '')
                                            if x_forma:
                                                item['FormaPago'] = x_forma
                                                item['forma_pago'] = x_forma
                                            if x_metodo:
                                                item['MetodoPago'] = x_metodo
                                                item['metodo_pago'] = x_metodo
                                        except Exception:
                                            pass
                                    except Exception:
                                        pass
                                    break
                            if xml_descargado:
                                break
                except Exception:
                    pass

                # 2. PDF Oficial
                try:
                    xpath_pdf = f"//table[@id='ctl00_MainContent_tblResult']//tr[contains(., '{folio}')]//span[contains(@id, 'BtnRI')] | //table[@id='ctl00_MainContent_tblResult']//tr[contains(., '{folio}')]//a[contains(@onclick, 'Recupera')]"
                    btn_pdf = driver.find_elements(By.XPATH, xpath_pdf)
                    if btn_pdf:
                        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", btn_pdf[0])
                        time.sleep(0.04)
                        driver.execute_script("arguments[0].click();", btn_pdf[0])
                        for _ in range(5):
                            time.sleep(0.2)
                            for fname in os.listdir(download_dir):
                                if fname.lower().endswith(".pdf") and not fname.lower().endswith(".tmp"):
                                    src_f = os.path.join(download_dir, fname)
                                    dst_f = os.path.join(pdf_dir, f"{folio}.pdf")
                                    try:
                                        shutil.move(src_f, dst_f)
                                        pdf_descargado = True
                                    except Exception:
                                        pass
                                    break
                            if pdf_descargado:
                                break

                        if len(driver.window_handles) > 1:
                            for h in driver.window_handles:
                                if h != main_win:
                                    try:
                                        driver.switch_to.window(h)
                                        driver.close()
                                    except Exception:
                                        pass
                            driver.switch_to.window(main_win)
                except Exception:
                    pass

    return registros, comprobantes_dir, len(registros)

def run_extraction(tipo, year, month, cer_path, key_path, password, rfc, output_excel_path, download_dir=None, download_files=True):
    tipo = tipo.lower().strip()
    if tipo not in ["emitidas", "recibidas", "ambas", "ambos"]:
        emit({"type": "error", "error": "El tipo debe ser 'emitidas', 'recibidas' o 'ambos'"})
        return False

    browser_exe = find_browser_executable()
    if not browser_exe:
        emit({"type": "error", "error": "No se encontró ningún navegador compatible (Brave, Chrome o Edge)."})
        return False

    is_ambos = tipo in ["ambas", "ambos"]
    nombre_tipo = "Facturas Emitidas e Ingresos + Recibidas y Gastos" if is_ambos else ("Facturas Emitidas (Clientes)" if tipo == "emitidas" else "Facturas Recibidas (Proveedores)")
    emit({"type": "log", "msg": f"ðŸš€ Iniciando motor de sincronización SAT para {nombre_tipo}..."})
    emit({"type": "log", "msg": f"ðŸ‘¤ Contribuyente RFC: {rfc}"})
    periodo_str = MESES_NOMBRES[int(month)] if str(month).isdigit() and 1 <= int(month) <= 12 else ('Mes ' + str(month) if str(month) != 'all' else 'Consolidado Anual')
    emit({"type": "log", "msg": f"📅 Ejercicio: {year} • Periodo: {periodo_str}"})
    emit({"type": "progress", "value": 5})

    temp_profile_dir = tempfile.mkdtemp(prefix="sat_facturas_")
    if not download_dir:
        download_dir = os.path.dirname(os.path.abspath(output_excel_path))
    os.makedirs(download_dir, exist_ok=True)

    options = Options()
    options.binary_location = browser_exe
    options.page_load_strategy = "eager"
    options.add_argument("--headless=new")
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-popup-blocking")
    options.add_argument("--window-size=1920,1080")
    options.add_argument("--disable-blink-features=AutomationControlled")
    options.add_argument("--ignore-certificate-errors")
    options.add_argument("--allow-running-insecure-content")
    options.add_argument(f"--user-data-dir={temp_profile_dir}")
    options.add_argument("--disable-cache")
    options.add_argument("--disk-cache-size=0")
    options.add_argument("--no-first-run")
    options.add_argument("--no-default-browser-check")
    options.add_argument("--disable-application-cache")
    options.add_argument("--disable-session-crashed-bubble")
    options.add_experimental_option("excludeSwitches", ["enable-automation"])
    options.add_experimental_option("useAutomationExtension", False)
    options.add_argument("user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36")

    prefs = {
        "download.default_directory": download_dir,
        "download.prompt_for_download": False,
        "download.directory_upgrade": True,
        "plugins.always_open_pdf_externally": True,
        "safebrowsing.enabled": True,
        "profile.default_content_settings.popups": 0,
        "profile.default_content_setting_values.cookies": 1,
        "profile.cookie_controls_mode": 0,
        "profile.block_third_party_cookies": False,
    }
    options.add_experimental_option("prefs", prefs)

    driver = None

    try:
        service = Service()
        if sys.platform == "win32":
            service.creation_flags = subprocess.CREATE_NO_WINDOW

        driver = webdriver.Chrome(service=service, options=options)
        driver.set_page_load_timeout(120)
        driver.set_script_timeout(120)

        try:
            driver.execute_cdp_cmd('Browser.setDownloadBehavior', {
                'behavior': 'allowAndName',
                'downloadPath': download_dir,
                'eventsEnabled': True
            })
        except Exception:
            pass
        driver.execute_cdp_cmd('Page.setDownloadBehavior', {
            'behavior': 'allow',
            'downloadPath': download_dir
        })

        emit({"type": "log", "msg": "ðŸŒ Conectando al portal de autenticación del SAT..."})
        emit({"type": "progress", "value": 12})
        try:
            driver.get(SAT_LOGIN_URL)
        except Exception:
            pass
        time.sleep(2)

        # Autenticación con e.firma
        emit({"type": "log", "msg": "ðŸ”‘ Autenticando ante el SAT con e.firma (.cer, .key)..."})
        emit({"type": "progress", "value": 20})

        curr_initial_url = driver.current_url.lower()
        ya_autenticado = ("portalcfdi.facturaelectronica.sat.gob.mx" in curr_initial_url and "cfdiau" not in curr_initial_url) or \
                         bool(driver.find_elements(By.ID, "ctl00_LblRfcAutenticado"))

        if not ya_autenticado:
            if not os.path.exists(cer_path) or not os.path.exists(key_path):
                emit({"type": "error", "error": "Archivos .cer o .key no encontrados en el servidor."})
                return False

            wait_login = WebDriverWait(driver, 20)

            if "satx509custom" not in driver.current_url.lower() and not driver.find_elements(By.ID, "fileCertificate"):
                try:
                    btn_fiel = wait_login.until(EC.element_to_be_clickable((By.ID, "buttonFiel")))
                    try:
                        btn_fiel.click()
                    except Exception:
                        driver.execute_script("arguments[0].click();", btn_fiel)
                    wait_login.until(lambda d: "satx509custom" in d.current_url.lower() or bool(d.find_elements(By.ID, "fileCertificate")))
                except Exception:
                    pass

            wait_login.until(EC.presence_of_element_located((By.ID, "fileCertificate")))
            wait_login.until(EC.presence_of_element_located((By.ID, "filePrivateKey")))
            wait_login.until(EC.presence_of_element_located((By.ID, "privateKeyPassword")))

            input_cer = driver.find_element(By.ID, "fileCertificate")
            input_key = driver.find_element(By.ID, "filePrivateKey")
            input_pass = driver.find_element(By.ID, "privateKeyPassword")

            input_cer.send_keys(cer_path)
            driver.execute_script("arguments[0].dispatchEvent(new Event('change', { bubbles: true }));", input_cer)
            time.sleep(0.5)

            input_key.send_keys(key_path)
            driver.execute_script("arguments[0].dispatchEvent(new Event('change', { bubbles: true }));", input_key)
            time.sleep(0.5)

            input_pass.clear()
            input_pass.send_keys(password)
            driver.execute_script("arguments[0].dispatchEvent(new Event('change', { bubbles: true }));", input_pass)
            time.sleep(0.5)

            for _ in range(15):
                c = driver.find_element(By.ID, "txtCertificate").get_attribute("value") or ""
                k = driver.find_element(By.ID, "txtPrivateKey").get_attribute("value") or ""
                r = driver.find_element(By.ID, "rfc").get_attribute("value") or ""
                if c and k and len(r) >= 10:
                    break
                time.sleep(0.5)

            btn_submit = driver.find_element(By.ID, "submit")
            emit({"type": "log", "msg": "⏳ Validando credenciales criptográficas de la e.firma..."})
            emit({"type": "progress", "value": 28})
            try:
                btn_submit.click()
            except Exception:
                driver.execute_script("arguments[0].click();", btn_submit)

            inicio_login = time.time()
            autenticado = False
            while time.time() - inicio_login < 45:
                time.sleep(0.8)

                curr = driver.current_url.lower()

                # Detectar error formal devuelto por el portal del SAT
                try:
                    err_box = driver.find_elements(By.CSS_SELECTOR, ".alert-danger, #divError, .error-message, .msgError")
                    for eb in err_box:
                        if eb.is_displayed() and eb.text.strip():
                            emit({"type": "error", "error": f"Error del SAT en login: {eb.text.strip()}"})
                            return False
                except Exception:
                    pass

                try:
                    page_text = driver.find_element(By.TAG_NAME, "body").text
                    if "Ha ocurrido un error al procesar su última acción" in page_text:
                        emit({"type": "error", "error": "Error del SAT en login: Ha ocurrido un error al procesar su última acción"})
                        return False
                    if "Acceso no autorizado" in page_text or "Contraseña incorrecta" in page_text or "Clave privada no corresponde" in page_text:
                        emit({"type": "error", "error": f"Error del SAT en login: {page_text[:120]}"})
                        return False
                except Exception:
                    pass

                if "portalcfdi.facturaelectronica.sat.gob.mx" in curr and "cfdiau" not in curr:
                    autenticado = True
                    break
                if driver.find_elements(By.ID, "ctl00_LblRfcAutenticado"):
                    autenticado = True
                    break
                try:
                    if "cfdiau.sat.gob.mx" not in driver.current_url.lower():
                        autenticado = True
                except Exception:
                    pass

            if not autenticado:
                emit({"type": "error", "error": "El servidor del SAT tardó demasiado en autenticar la e.firma."})
                return False

        emit({"type": "log", "msg": "✅ Autenticación exitosa. Sesión oficial abierta."})
        emit({"type": "progress", "value": 35})

        int_year = int(year)
        is_all_months = str(month).lower() == "all"

        if is_ambos:
            # Extracción combinada en una sola sesión: Emitidas y Recibidas
            emit({"type": "log", "msg": "ðŸ“¦ [1/2] Extrayendo comprobantes EMITIDOS (Ingresos)..."})
            emitidas, comp_dir_em, total_em = extraer_modulo_facturas(
                driver, "emitidas", year, month, rfc, download_files, download_dir, emit, progress_range=(35, 65)
            )

            emit({"type": "log", "msg": "ðŸ“¦ [2/2] Extrayendo comprobantes RECIBIDOS (Gastos)..."})
            recibidas, comp_dir_rec, total_rec = extraer_modulo_facturas(
                driver, "recibidas", year, month, rfc, download_files, download_dir, emit, progress_range=(65, 95)
            )

            output_excel = None
            if download_files:
                output_excel = generar_excel_facturas(emitidas + recibidas, "ambos", output_excel_path)

            emit({"type": "progress", "value": 100})
            emit({
                "type": "finished",
                "data": {
                    "tipo": "ambos",
                    "year": str(int_year),
                    "month": "all" if is_all_months else str(month),
                    "total_registros": len(emitidas) + len(recibidas),
                    "total_emitidas": len(emitidas),
                    "total_recibidas": len(recibidas),
                    "emitidas": emitidas,
                    "recibidas": recibidas,
                    "registros": emitidas + recibidas,
                    "excel_path": output_excel
                }
            })
            return True

        else:
            # Extracción simple (emitidas o recibidas)
            registros, comprobantes_dir, total_registros = extraer_modulo_facturas(
                driver, tipo, year, month, rfc, download_files, download_dir, emit, progress_range=(35, 88)
            )

            output_excel = None
            if download_files:
                output_excel = generar_excel_facturas(registros, tipo, output_excel_path)

            zip_filepath = None
            zip_filename = None
            total_archivos = 0

            if download_files and comprobantes_dir and os.path.exists(comprobantes_dir):
                monthTag = "ANUAL" if is_all_months else f"M{month}"
                zip_filename = f"Facturas_{tipo.upper()}_{rfc}_{int_year}_{monthTag}_PDF_XML.zip"
                zip_filepath = os.path.join(download_dir, zip_filename)

                emit({"type": "progress", "value": 94})
                emit({"type": "log", "msg": "ðŸ“¦ Comprimiendo comprobantes fiscales (PDF y XML) en archivo ZIP de auditoría..."})

                try:
                    for item in os.listdir(comprobantes_dir):
                        item_path = os.path.join(comprobantes_dir, item)
                        if os.path.isdir(item_path):
                            has_files = any(len(fls) > 0 for _, _, fls in os.walk(item_path))
                            if not has_files:
                                shutil.rmtree(item_path, ignore_errors=True)

                    compression = zipfile.ZIP_DEFLATED if hasattr(zipfile, 'ZIP_DEFLATED') else zipfile.ZIP_STORED
                    with zipfile.ZipFile(zip_filepath, 'w', compression=compression) as zf:
                        for root, _, files in os.walk(comprobantes_dir):
                            for f in files:
                                full_p = os.path.join(root, f)
                                rel_p = os.path.relpath(full_p, comprobantes_dir)
                                zf.write(full_p, rel_p)
                                total_archivos += 1
                    emit({"type": "log", "msg": f"✅ Archivo ZIP generado exitosamente con {total_archivos} comprobantes (PDF + XML)."})
                except Exception as e:
                    emit({"type": "log", "msg": f"⚠️ï¸ Aviso al generar ZIP: {str(e)}"})

            total_monto = round(sum(r.get("Total", 0.0) for r in registros), 2)
            vigentes_count = sum(1 for r in registros if "vigente" in (r.get("Estado del Comprobante") or "").lower())
            cancelados_count = sum(1 for r in registros if "cancelado" in (r.get("Estado del Comprobante") or "").lower())

            emit({"type": "progress", "value": 100})
            emit({
                "type": "finished",
                "data": {
                    "tipo": tipo,
                    "year": str(int_year),
                    "month": "all" if is_all_months else str(month),
                    "total_registros": len(registros),
                    "total_monto": total_monto,
                    "vigentes": vigentes_count,
                    "cancelados": cancelados_count,
                    "registros": registros,
                    "excel_path": output_excel,
                    "zip_path": zip_filepath if (zip_filepath and os.path.exists(zip_filepath) and total_archivos > 0) else None,
                    "zip_filename": zip_filename,
                    "archivos_count": total_archivos,
                    "comprobantes_dir": comprobantes_dir
                }
            })
            return True

    except Exception as e:
        err_msg = str(e)
        emit({"type": "error", "error": f"Error durante la extracción del SAT: {err_msg[:250]}"})
        return False
    finally:
        if driver:
            try:
                driver.delete_all_cookies()
            except Exception:
                pass
            try:
                for h in driver.window_handles:
                    try:
                        driver.switch_to.window(h)
                        driver.close()
                    except Exception:
                        pass
            except Exception:
                pass
            try:
                driver.quit()
            except Exception:
                pass
            driver = None

        if temp_profile_dir and os.path.exists(temp_profile_dir):
            try:
                time.sleep(0.5)
                shutil.rmtree(temp_profile_dir, ignore_errors=True)
            except Exception:
                pass

def main():
    parser = argparse.ArgumentParser(description="Extractor Automatizado de Facturas SAT (Emitidas/Recibidas)")
    parser.add_argument("--config", help="Ruta a archivo config.json con los parámetros")
    parser.add_argument("--tipo", choices=["emitidas", "recibidas", "ambas", "ambos"], default="emitidas")
    parser.add_argument("--year", default=str(datetime.now().year))
    parser.add_argument("--month", default=str(datetime.now().month))
    parser.add_argument("--cer", help="Ruta a .cer")
    parser.add_argument("--key", help="Ruta a .key")
    parser.add_argument("--password", help="Contraseña de e.firma")
    parser.add_argument("--rfc", help="RFC del contribuyente")
    parser.add_argument("--output", help="Ruta de salida del archivo Excel")
    parser.add_argument("--download-dir", help="Directorio para descargas")

    parser.add_argument("--download-files", action="store_true", default=None, help="Si se deben descargar los archivos PDF y XML")
    parser.add_argument("--fast-sync", action="store_true", default=False, help="Modo rápido solo de metadatos fiscales")

    args = parser.parse_args()

    if args.config and os.path.exists(args.config):
        with open(args.config, "r", encoding="utf-8") as f:
            cfg = json.load(f)
        tipo = cfg.get("tipo", "emitidas")
        year = cfg.get("year", str(datetime.now().year))
        month = cfg.get("month", str(datetime.now().month))
        cer_path = cfg.get("cer_path")
        key_path = cfg.get("key_path")
        password = cfg.get("password")
        rfc = cfg.get("rfc")
        output_path = cfg.get("output_excel")
        download_dir = cfg.get("download_dir")
        fast_sync = cfg.get("fast_sync", False)
        download_files = cfg.get("download_files", not fast_sync)
    else:
        tipo = args.tipo
        year = args.year
        month = args.month
        cer_path = args.cer
        key_path = args.key
        password = args.password
        rfc = args.rfc
        output_path = args.output
        download_dir = args.download_dir
        fast_sync = args.fast_sync
        download_files = args.download_files if args.download_files is not None else not fast_sync

    if not output_path:
        output_path = os.path.abspath(f"Reporte_Facturas_{tipo.capitalize()}_{rfc}_{year}_{month}.xlsx")

    success = run_extraction(
        tipo=tipo,
        year=year,
        month=month,
        cer_path=cer_path,
        key_path=key_path,
        password=password,
        rfc=rfc,
        output_excel_path=output_path,
        download_dir=download_dir,
        download_files=download_files
    )

    sys.exit(0 if success else 1)

if __name__ == "__main__":
    main()
