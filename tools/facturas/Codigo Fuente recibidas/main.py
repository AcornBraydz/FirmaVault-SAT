import tkinter as tk
from tkinter import filedialog, messagebox  # Para ventanas emergentes
from pdfminer.high_level import extract_text
import pandas as pd  # Para manejar los datos en Excel
from openpyxl import load_workbook  # Para aplicar las fórmulas y ajustes
from openpyxl.styles import Alignment  # Para centrar el contenido
from openpyxl.utils import get_column_letter
from datetime import datetime  # Para manejar y formatear fechas

def select_and_process_pdf():
    # Abrir cuadro de diálogo para seleccionar un archivo PDF
    file_path = filedialog.askopenfilename(
        title="Selecciona un archivo PDF",
        filetypes=[("Archivos PDF", "*.pdf")]
    )
    if file_path:
        process_pdf_to_excel(file_path)

def format_fecha_emision(fecha_emision):
    try:
        # Interpretar el formato original de la fecha
        fecha_original = datetime.strptime(fecha_emision, "%Y-%m-%dT%H")
    except ValueError:
        # Si el formato no coincide, devolver la fecha sin cambios
        return fecha_emision

    # Formatear la fecha al estilo deseado (DD-Mes-YYYY)
    fecha_formateada = fecha_original.strftime("%d-%B-%Y").lower()

    # Traducir el nombre del mes al español
    meses = {
        "january": "enero", "february": "febrero", "march": "marzo",
        "april": "abril", "may": "mayo", "june": "junio",
        "july": "julio", "august": "agosto", "september": "septiembre",
        "october": "octubre", "november": "noviembre", "december": "diciembre"
    }
    for en, es in meses.items():
        fecha_formateada = fecha_formateada.replace(en, es)

    return fecha_formateada.capitalize()

def process_pdf_to_excel(pdf_path):
    try:
        # Extraer todo el texto del archivo PDF
        text = extract_text(pdf_path)
        registros = []
        lines = text.split("\n")

        # Procesar las líneas extraídas del PDF
        i = 0
        while i < len(lines):
            if "Folio Fiscal:" in lines[i]:
                folio = lines[i].split(":")[1].strip()
                nombre = ""
                fecha_emision = ""
                rfc_emisor = ""
                estado_comprobante = ""
                efecto_comprobante = ""
                total = 0.00

                # Buscar "RFC Emisor"
                for j in range(i + 1, len(lines)):
                    if "RFC Emisor:" in lines[j]:
                        rfc_emisor = lines[j].split(":")[1].strip()
                        break

                # Buscar "Nombre o Razón Social"
                for j in range(i + 1, len(lines)):
                    if "Nombre o Razón Social:" in lines[j]:
                        nombre = lines[j].split(":")[1].strip()
                        for k in range(j + 1, len(lines)):
                            if lines[k].strip() and not ("RFC Receptor:" in lines[k] or "Folio Fiscal:" in lines[k]):
                                nombre += " " + lines[k].strip()
                            else:
                                break
                        break

                # Buscar "Fecha Emisión"
                for j in range(i + 1, len(lines)):
                    if "Fecha Emisión:" in lines[j]:
                        fecha_emision_raw = lines[j].split(":")[1].strip()
                        fecha_emision = format_fecha_emision(fecha_emision_raw)
                        break

                # Buscar "Total:"
                for j in range(i + 1, len(lines)):
                    if "Total:" in lines[j]:
                        total_text = lines[j].split(":")[1].strip()
                        total = float(total_text.replace("$", "").replace(",", ""))
                        break

                # Buscar "Estado del Comprobante" y "Efecto del Comprobante"
                for j in range(i + 1, len(lines)):
                    if "Estado del Comprobante:" in lines[j]:
                        partes = lines[j].split("Efecto del Comprobante:")
                        estado_comprobante = partes[0].split(":")[1].strip()
                        if len(partes) > 1:
                            efecto_comprobante = partes[1].strip()
                        break

                # Crear un registro como diccionario
                registros.append({
                    "Folio": folio,
                    "Fecha de Emisión": fecha_emision,
                    "Nombre o Razón Social": nombre,
                    "RFC": rfc_emisor,
                    "Importe": "",  # Se calculará con fórmula
                    "IVA al 16%": "",  # Se calculará con fórmula
                    "Total": total,
                    "Estado del Comprobante": estado_comprobante,
                    "Efecto del Comprobante": efecto_comprobante
                })
            i += 1

        # Convertir a DataFrame y exportar
        df = pd.DataFrame(registros)
        output_excel_path = pdf_path.replace(".pdf", "_procesado.xlsx")
        df.to_excel(output_excel_path, index=False)

        # Aplicar fórmulas y formato
        apply_formulas_and_formatting(output_excel_path)

        # Mostrar mensaje al usuario
        messagebox.showinfo("Procesamiento Completo", f"El archivo Excel ha sido generado con éxito:\n{output_excel_path}")

    except Exception as e:
        messagebox.showerror("Error", f"Ocurrió un error al procesar el archivo PDF:\n{e}")

def apply_formulas_and_formatting(output_excel_path):
    # Cargar el archivo Excel generado
    workbook = load_workbook(output_excel_path)
    sheet = workbook.active

    # Aplicar fórmula para "Importe" (columna E) = ROUND("Total" (columna G) / 1.16, 2)
    for row in range(2, sheet.max_row + 1):
        sheet[f"E{row}"] = f"=ROUND(G{row}/1.16, 2)"

    # Aplicar fórmula para "IVA al 16%" (columna F) = ROUND("Importe" (columna E) * 0.16, 2)
    for row in range(2, sheet.max_row + 1):
        sheet[f"F{row}"] = f"=ROUND(E{row}*0.16, 2)"

    # Ajustar ancho de las columnas y centrar contenido
    for column in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
            cell.alignment = Alignment(horizontal='center', vertical='center')
        sheet.column_dimensions[column_letter].width = max_length * 1.30 + 2

    # Guardar cambios
    workbook.save(output_excel_path)

def create_menu():
    root = tk.Tk()
    root.title("Procesador de PDF a Excel")
    root.geometry("400x200")

    label = tk.Label(root, text="Selecciona un archivo PDF", font=("Arial", 12))
    label.pack(pady=20)

    button = tk.Button(root, text="Seleccionar PDF", command=select_and_process_pdf, font=("Arial", 12), bg="lightblue")
    button.pack(pady=20)

    root.mainloop()

# Ejecutar el programa
create_menu()