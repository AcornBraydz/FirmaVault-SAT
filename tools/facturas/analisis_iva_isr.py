import sys
import os

# Configurar UTF-8 para stdout y stderr en Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
        sys.stderr.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

import json
import time
import argparse
from defusedxml import ElementTree as ET
from datetime import datetime
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Integración del Validador de Listas Negras EFOS (Art. 69-B CFF)
try:
    from sat_efos_validator import check_rfc_efos
except ImportError:
    try:
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        from sat_efos_validator import check_rfc_efos
    except Exception:
        def check_rfc_efos(rfc): return None

# Integración del Catálogo Maestro y Motor de Regímenes del SAT (19 regímenes oficiales)
try:
    from sat_regimenes import (
        CATALOGO_REGIMENES_SAT, normalizar_regimen_code,
        get_regimen_info, calcular_impuestos_regimen
    )
except ImportError:
    try:
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        from sat_regimenes import (
            CATALOGO_REGIMENES_SAT, normalizar_regimen_code,
            get_regimen_info, calcular_impuestos_regimen
        )
    except Exception:
        CATALOGO_REGIMENES_SAT = {}
        def normalizar_regimen_code(k): return str(k or "612")
        def get_regimen_info(k): return {"code": "612", "nombre": "General", "short_label": "612", "columna_header": "Utilidad Fiscal ISR", "fundamento": "LISR"}
        def calcular_impuestos_regimen(r, i, g, ri=0.0, riv=0.0):
            return {"regimen_code": "612", "base_isr": round(i - g, 2), "isr_determinado": 0.0, "isr_neto": 0.0, "columna_header": "Utilidad Fiscal ISR"}

def calcular_resico_mes(ingresos_cobrados, ret_isr=0.0):
    """
    Tabla mensual progresiva de RESICO Personas Físicas (Art. 113-E LISR):
    Hasta $25,000.00         -> 1.00%
    Hasta $50,000.00         -> 1.10%
    Hasta $83,333.33         -> 1.50%
    Hasta $208,333.33        -> 2.00%
    Hasta $291,666.66        -> 2.50%
    Más de $291,666.66 (hasta 3.5 MDP anual) -> 2.50%
    Aplica retención del 1.25% de personas morales (Art. 113-J LISR).
    """
    ing = max(0.0, float(ingresos_cobrados or 0.0))
    ret_isr = max(0.0, float(ret_isr or 0.0))
    if ing <= 0:
        return {
            "base": 0.0,
            "tasa": 0.0,
            "tasa_porcentaje": "0.0%",
            "isr_causado": 0.0,
            "retencion_125": round(ret_isr, 2),
            "isr_neto": 0.0,
            "isr_a_favor": round(ret_isr, 2)
        }
    if ing <= 25000.00:
        tasa = 0.0100
    elif ing <= 50000.00:
        tasa = 0.0110
    elif ing <= 83333.33:
        tasa = 0.0150
    elif ing <= 208333.33:
        tasa = 0.0200
    else:
        tasa = 0.0250

    isr_causado = round(ing * tasa, 2)
    isr_neto = max(0.0, round(isr_causado - ret_isr, 2))
    isr_a_favor = round(max(0.0, ret_isr - isr_causado), 2)
    return {
        "base": round(ing, 2),
        "tasa": tasa,
        "tasa_porcentaje": f"{round(tasa * 100, 2)}%",
        "isr_causado": isr_causado,
        "retencion_125": round(ret_isr, 2),
        "isr_neto": isr_neto,
        "isr_a_favor": isr_a_favor
    }

MESES_NOMBRES = [
    "", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"
]

CATALOGO_FORMAS_PAGO = {
    '01': 'Efectivo (01)',
    '02': 'Cheque nominativo (02)',
    '03': 'Transferencia electrónica (03)',
    '04': 'Tarjeta de crédito (04)',
    '05': 'Monedero electrónico (05)',
    '06': 'Dinero electrónico (06)',
    '08': 'Vales de despensa (08)',
    '12': 'Dación en pago (12)',
    '13': 'Pago por subrogación (13)',
    '14': 'Pago por consignación (14)',
    '15': 'Condonación (15)',
    '17': 'Compensación (17)',
    '23': 'Novación (23)',
    '24': 'Confusión (24)',
    '25': 'Remisión de deuda (25)',
    '26': 'Prescripción o caducidad (26)',
    '27': 'A satisfacción del acreedor (27)',
    '28': 'Tarjeta de débito (28)',
    '29': 'Tarjeta de servicios (29)',
    '30': 'Aplicación de anticipos (30)',
    '31': 'Intermediario pagos (31)',
    '99': 'Por definir (99)',
}

def get_forma_pago_nombre(code, metodo='PUE'):
    m = str(metodo or 'PUE').upper().strip()
    if not code or code in ['SIN_XML', 'sin_xml', 'SD', 'sd', '']:
        if 'REP' in m:
            return 'Bancarizado (REP)'
        if 'EGRESO' in m:
            return 'Nota de Crédito'
        if 'NOMINA' in m or 'NÓMINA' in m:
            return 'Nómina (99)'
        return 'Contado (PUE)' if m == 'PUE' else 'Crédito (PPD)'
    if code in ['REP', 'rep']:
        return 'Bancarizado (REP)'
    if code in ['EGRESO', 'egreso']:
        return 'Nota de Crédito'
    if code == '99':
        return 'Por definir (99)'
    c = str(code).strip().zfill(2)
    return CATALOGO_FORMAS_PAGO.get(c, f'Forma {code}')


def emit_event(data):
    """Emite evento estructurado en formato JSON a stdout para consumo en Node.js"""
    try:
        print(json.dumps(data, ensure_ascii=False), flush=True)
    except Exception:
        print(json.dumps(data, ensure_ascii=True), flush=True)

def parse_cfdi_xml(xml_content):
    """
    Parsea un archivo XML de CFDI 3.3 o 4.0 y extrae los atributos y nodos fiscales clave.
    """
    try:
        root = ET.fromstring(xml_content)
    except Exception as e:
        return None

    # Espacios de nombres del SAT
    ns = {
        'cfdi': 'http://www.sat.gob.mx/cfd/4',
        'cfdi3': 'http://www.sat.gob.mx/cfd/3',
        'pago20': 'http://www.sat.gob.mx/Pagos20',
        'pago10': 'http://www.sat.gob.mx/Pagos',
    }

    # Atributos raíz de Comprobante
    tipo_comp = root.attrib.get('TipoDeComprobante', 'I')
    fecha_raw = root.attrib.get('Fecha', '')
    subtotal = float(root.attrib.get('SubTotal', 0.0))
    descuento = float(root.attrib.get('Descuento', 0.0))
    total = float(root.attrib.get('Total', 0.0))
    metodo_pago = root.attrib.get('MetodoPago', 'PUE' if tipo_comp != 'P' else 'REP')
    forma_pago = root.attrib.get('FormaPago', '')
    folio_fiscal = ""

    is_gasolina = False
    doctos_relacionados = []
    cfdi_relacionados = []

    # Complemento TimbreFiscalDigital (UUID)
    for elem in root.iter():
        tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
        if tag == 'TimbreFiscalDigital':
            folio_fiscal = elem.attrib.get('UUID', '')
        elif tag == 'CfdiRelacionado':
            rel_u = elem.attrib.get('UUID', '').upper().strip()
            if rel_u and rel_u not in cfdi_relacionados:
                cfdi_relacionados.append(rel_u)
        elif tag == 'Pago':
            if not forma_pago and 'FormaDePagoP' in elem.attrib:
                forma_pago = elem.attrib.get('FormaDePagoP', '')
            if tipo_comp == 'P' and 'Monto' in elem.attrib and total == 0.0:
                try: total += float(elem.attrib.get('Monto', 0.0))
                except: pass
        elif tag == 'Concepto':
            clave = elem.attrib.get('ClaveProdServ', '')
            desc = elem.attrib.get('Descripcion', '').lower()
            if clave.startswith('151015') or any(k in desc for k in ['gasolina', 'combustible', 'diesel', 'magna', 'premium']):
                is_gasolina = True
        elif tag == 'DoctoRelacionado' and tipo_comp == 'P':
            dr_uuid = elem.attrib.get('IdDocumento', '')
            imp_pagado = float(elem.attrib.get('ImpPagado', 0.0)) if 'ImpPagado' in elem.attrib else 0.0
            iva_16_dr = 0.0
            iva_8_dr = 0.0
            for dr_child in elem.iter():
                dr_tag = dr_child.tag.split('}')[-1] if '}' in dr_child.tag else dr_child.tag
                if dr_tag == 'TrasladoDR':
                    imp = dr_child.attrib.get('ImpuestoDR', '')
                    tasa = float(dr_child.attrib.get('TasaOCuotaDR', 0.0)) if 'TasaOCuotaDR' in dr_child.attrib else 0.0
                    imp_traslado = float(dr_child.attrib.get('ImporteDR', 0.0)) if 'ImporteDR' in dr_child.attrib else 0.0
                    if imp in ['002', 'IVA']:
                        if abs(tasa - 0.16) < 0.005: iva_16_dr += imp_traslado
                        elif abs(tasa - 0.08) < 0.005: iva_8_dr += imp_traslado
            doctos_relacionados.append({'uuid': dr_uuid, 'imp_pagado': imp_pagado, 'iva_16': iva_16_dr, 'iva_8': iva_8_dr})

    if not forma_pago:
        forma_pago = '99' if tipo_comp != 'P' else '03'
    forma_pago = str(forma_pago).strip().zfill(2)

    # Emisor y Receptor
    emisor_rfc, emisor_nombre = "", ""
    receptor_rfc, receptor_nombre = "", ""

    for elem in root:
        tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
        if tag == 'Emisor':
            emisor_rfc = elem.attrib.get('Rfc', '')
            emisor_nombre = elem.attrib.get('Nombre', emisor_rfc)
        elif tag == 'Receptor':
            receptor_rfc = elem.attrib.get('Rfc', '')
            receptor_nombre = elem.attrib.get('Nombre', receptor_rfc)

    # Impuestos Globales
    iva_16 = 0.0
    iva_8 = 0.0
    iva_0 = 0.0
    iva_exento = 0.0
    ret_isr = 0.0
    ret_iva = 0.0

    # Recorrer nodos de Impuestos o Conceptos
    for elem in root.iter():
        tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
        if tag == 'Traslado':
            impuesto = elem.attrib.get('Impuesto', '')
            tasa = float(elem.attrib.get('TasaOCuota', 0.0)) if 'TasaOCuota' in elem.attrib else 0.0
            importe = float(elem.attrib.get('Importe', 0.0)) if 'Importe' in elem.attrib else 0.0
            tipo_factor = elem.attrib.get('TipoFactor', '')

            if impuesto in ['002', 'IVA']:
                if tipo_factor == 'Exento':
                    iva_exento += subtotal
                elif abs(tasa - 0.16) < 0.005:
                    iva_16 += importe
                elif abs(tasa - 0.08) < 0.005:
                    iva_8 += importe
                elif abs(tasa) < 0.001:
                    iva_0 += importe

        elif tag == 'Retencion':
            impuesto = elem.attrib.get('Impuesto', '')
            importe = float(elem.attrib.get('Importe', 0.0)) if 'Importe' in elem.attrib else 0.0
            if impuesto in ['001', 'ISR']:
                ret_isr += importe
            elif impuesto in ['002', 'IVA']:
                ret_iva += importe

    dt = None
    fecha_display = fecha_raw
    if fecha_raw:
        try:
            dt = datetime.fromisoformat(fecha_raw.split("T")[0])
            if 1 <= dt.month <= 12:
                fecha_display = f"{dt.day:02d}-{MESES_NOMBRES[dt.month]}-{dt.year}"
        except Exception:
            pass

    # Evaluación fiscal y listas negras SAT (EFOS Art. 69-B CFF & Reglas LISR)
    efos_info = check_rfc_efos(emisor_rfc)
    is_efos_definitivo = bool(efos_info and efos_info.get('es_definitivo'))

    es_deducible = True
    no_deducible_tipo = ""
    riesgo_deducibilidad = ""

    if is_efos_definitivo:
        es_deducible = False
        no_deducible_tipo = "efos_definitivo"
        riesgo_deducibilidad = f"🚨 EFOS Definitivo (Art. 69-B CFF): Factura No Deducible"
    elif is_gasolina and forma_pago == '01':
        es_deducible = False
        no_deducible_tipo = "gasolina_efectivo"
        riesgo_deducibilidad = "⚠️ï¸ Gasolina en Efectivo: 100% No Deducible (Art. 27 Fracc. III LISR)"
    elif forma_pago == '01' and total > 2000.0:
        es_deducible = False
        no_deducible_tipo = "efectivo_mayor_2k"
        riesgo_deducibilidad = "⚠️ï¸ Pago en Efectivo > $2,000: No Deducible (Art. 27 Fracc. III LISR)"
    elif tipo_comp == 'N':
        es_deducible = True
        riesgo_deducibilidad = "ℹ️ï¸ CFDI de Nómina (Sueldos y Salarios - IVA $0.00)"

    subtotal_deducible = subtotal if es_deducible else 0.0
    iva_16_acreditable = iva_16 if es_deducible else 0.0
    iva_8_acreditable = iva_8 if es_deducible else 0.0

    return {
        'uuid': folio_fiscal,
        'fecha': fecha_display,
        'fecha_raw': fecha_raw,
        'mes': dt.month if dt else 1,
        'anio': dt.year if dt else 2026,
        'tipo': tipo_comp,
        'metodo_pago': metodo_pago,
        'forma_pago': forma_pago,
        'emisor_rfc': emisor_rfc,
        'emisor_nombre': emisor_nombre,
        'receptor_rfc': receptor_rfc,
        'receptor_nombre': receptor_nombre,
        'subtotal': subtotal,
        'subtotal_deducible': subtotal_deducible,
        'descuento': descuento,
        'iva_16': round(iva_16_acreditable, 2),
        'iva_16_original': round(iva_16, 2),
        'iva_8': round(iva_8_acreditable, 2),
        'iva_8_original': round(iva_8, 2),
        'iva_0': round(iva_0, 2),
        'iva_exento': round(iva_exento, 2),
        'ret_isr': round(ret_isr, 2),
        'ret_iva': round(ret_iva, 2),
        'total': round(total, 2),
        'is_gasolina': is_gasolina,
        'es_deducible': es_deducible,
        'no_deducible_tipo': no_deducible_tipo,
        'efos_info': efos_info,
        'riesgo_deducibilidad': riesgo_deducibilidad,
        'doctos_relacionados': doctos_relacionados,
        'relacionados': cfdi_relacionados
    }

def generar_papel_trabajo_excel(datos_emitidas, datos_recibidas, year, zone, output_path, report_type='completo', month=0, regime='612', opcion_arrendamiento='ciega'):
    """
    Genera el libro Excel profesional con 5 pestañas y fórmulas contables vivas.
    """
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    # Paletas de color corporativas
    COLOR_NAVY = "0F172A"       # Headers Resumen
    COLOR_INDIGO = "312E81"     # Headers Ingresos
    COLOR_TEAL = "134E4A"       # Headers Gastos
    COLOR_AMBER = "78350F"      # Headers Retenciones

    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_title = Font(name="Calibri", size=14, bold=True, color=COLOR_NAVY)

    fill_navy = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
    fill_indigo = PatternFill(start_color=COLOR_INDIGO, end_color=COLOR_INDIGO, fill_type="solid")
    fill_teal = PatternFill(start_color=COLOR_TEAL, end_color=COLOR_TEAL, fill_type="solid")
    fill_amber = PatternFill(start_color=COLOR_AMBER, end_color=COLOR_AMBER, fill_type="solid")
    fill_total = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    total_border = Border(
        top=Side(style='thin', color='0F172A'),
        bottom=Side(style='double', color='0F172A')
    )

    num_format = "$#,##0.00"

    # ==========================================
    # PESTAÑA 1: Resumen Ejecutivo Anual
    # ==========================================
    ws_resumen = wb.create_sheet(title="1_Resumen_Ejecutivo")
    ws_resumen.views.sheetView[0].showGridLines = True

    regime_label = "RESICO" if regime == "resico" else ("Persona Moral" if regime == "persona_moral" else "Actividad Empresarial")
    periodo_title = f"{MESES_NOMBRES[month].upper()} {year}" if (month and month > 0) else f"EJERCICIO {year}"
    ws_resumen["A1"] = f"ANÁLISIS DE IVA E ISR ({regime_label.upper()}) - {periodo_title}"
    ws_resumen["A1"].font = font_title
    zona_str = "Zona General (16% CDMX / Edoméx)" if zone == 'general' else "Zona Fronteriza (8% Estímulo)"
    ws_resumen["A2"] = f"Configuración: {regime_label} | {zona_str} | Determinación Mensual por Flujo de Efectivo"
    ws_resumen["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

    # Tarjetas KPI de Indicadores Clave
    fill_kpi = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    font_kpi_label = Font(name="Calibri", size=9, bold=True, color="475569")
    font_kpi_val = Font(name="Calibri", size=13, bold=True, color="0F172A")
    kpi_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    meses_a_procesar = [month] if (month and month > 0) else list(range(1, 13))
    tot_row = 9 + len(meses_a_procesar)
    kpi_suffix = f" ({MESES_NOMBRES[month].upper()})" if (month and month > 0) else " (ANUAL)"

    ws_resumen.merge_cells("A4:C4")
    ws_resumen.merge_cells("A5:C5")
    ws_resumen["A4"] = f"TOTAL INGRESOS EMITIDOS{kpi_suffix}"
    ws_resumen["A5"] = f"=B{tot_row}"

    ws_resumen.merge_cells("D4:F4")
    ws_resumen.merge_cells("D5:F5")
    ws_resumen["D4"] = f"TOTAL GASTOS DEDUCIBLES{kpi_suffix}"
    ws_resumen["D5"] = f"=E{tot_row}"

    ws_resumen.merge_cells("G4:I4")
    ws_resumen.merge_cells("G5:I5")
    ws_resumen["G4"] = f"DIFERENCIA / SALDO NETO IVA{kpi_suffix}"
    ws_resumen["G5"] = f"=I{tot_row}"

    # Normalizar régimen oficial del SAT
    reg_code = normalizar_regimen_code(regime)
    reg_info = get_regimen_info(reg_code)
    col_header_isr = reg_info.get("columna_header", "Utilidad Fiscal ISR")

    ws_resumen.merge_cells("J4:K4")
    ws_resumen.merge_cells("J5:K5")
    if reg_code == '626':
        ws_resumen["J4"] = f"ESTIMADO ISR RESICO (NETO){kpi_suffix}"
    elif reg_code == '606':
        arrend_tag = "CIEGA 35%" if opcion_arrendamiento == "ciega" else "GASTOS REALES"
        ws_resumen["J4"] = f"UTILIDAD ARRENDAMIENTO ({arrend_tag}){kpi_suffix}"
    elif reg_code == 'resico_pm':
        ws_resumen["J4"] = f"UTILIDAD FISCAL RESICO PM{kpi_suffix}"
    else:
        ws_resumen["J4"] = f"{col_header_isr.upper()}{kpi_suffix}"
    ws_resumen["J5"] = f"=K{tot_row}"

    for r_idx in [4, 5]:
        for c_idx in range(1, 12):
            cell = ws_resumen.cell(row=r_idx, column=c_idx)
            cell.fill = fill_kpi
            cell.border = kpi_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if r_idx == 4:
                cell.font = font_kpi_label
            else:
                cell.font = font_kpi_val
                cell.number_format = num_format

    ret_isr_header = "Retención ISR 1.25% (Art. 113-J)" if reg_code == '626' else "Retención ISR Sufrida"
    headers_resumen = [
        "Mes", "Subtotal Emitidas", "IVA 16% Emitido", "IVA 8% Emitido",
        "Subtotal Recibidas Deducibles", "IVA 16% Acreditable", "IVA 8% Acreditable",
        "Retención IVA Sufrida", "Diferencia IVA (A Cargo / A Favor)",
        ret_isr_header,
        col_header_isr
    ]

    for col_idx in range(1, len(headers_resumen) + 1):
        cell = ws_resumen.cell(row=8, column=col_idx)
        cell.value = headers_resumen[col_idx - 1]
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    curr_row = 9
    for m in meses_a_procesar:
        row_num = curr_row
        em_mes = [x for x in datos_emitidas if x.get('mes') == m]
        rec_mes = [x for x in datos_recibidas if x.get('mes') == m]

        sub_em = sum(x['subtotal'] for x in em_mes)
        iva16_em = sum(x['iva_16'] for x in em_mes)
        iva8_em = sum(x['iva_8'] for x in em_mes)

        # Filtrar exclusivamente gastos deducibles e IVA acreditable
        sub_rec = sum(x.get('subtotal_deducible', x['subtotal']) for x in rec_mes if x.get('es_deducible', True))
        iva16_rec = sum(x['iva_16'] for x in rec_mes if x.get('es_deducible', True))
        iva8_rec = sum(x['iva_8'] for x in rec_mes if x.get('es_deducible', True))

        ret_iva_suf = sum(x['ret_iva'] for x in em_mes)
        ret_isr_suf = sum(x['ret_isr'] for x in em_mes)

        ws_resumen.cell(row=row_num, column=1, value=MESES_NOMBRES[m])
        ws_resumen.cell(row=row_num, column=2, value=sub_em)
        ws_resumen.cell(row=row_num, column=3, value=iva16_em)
        ws_resumen.cell(row=row_num, column=4, value=iva8_em)
        ws_resumen.cell(row=row_num, column=5, value=sub_rec)
        ws_resumen.cell(row=row_num, column=6, value=iva16_rec)
        ws_resumen.cell(row=row_num, column=7, value=iva8_rec)
        ws_resumen.cell(row=row_num, column=8, value=ret_iva_suf)
        ws_resumen.cell(row=row_num, column=9, value=f"=(C{row_num}+D{row_num})-(F{row_num}+G{row_num})-H{row_num}")
        ws_resumen.cell(row=row_num, column=10, value=ret_isr_suf)

        calc_reg_m = calcular_impuestos_regimen(
            reg_code, sub_em, sub_rec,
            retencion_isr=ret_isr_suf, retencion_iva=ret_iva_suf,
            opcion_arrendamiento=opcion_arrendamiento
        )
        if reg_code == '626':
            # Fórmula viva LISR Art. 113-E: Tarifa RESICO PF sobre ingresos cobrados menos Retención 1.25%
            ws_resumen.cell(
                row=row_num, column=11,
                value=f"=MAX(0, (B{row_num}*IF(B{row_num}<=25000, 0.01, IF(B{row_num}<=50000, 0.011, IF(B{row_num}<=83333.33, 0.015, IF(B{row_num}<=208333.33, 0.02, 0.025))))) - J{row_num})"
            )
        elif reg_code == '606':
            if opcion_arrendamiento == 'ciega':
                # Fórmula viva Art. 115 LISR: Ingresos menos 35% de deducción ciega
                ws_resumen.cell(row=row_num, column=11, value=f"=B{row_num}-(B{row_num}*0.35)")
            else:
                # Fórmula viva Art. 115 LISR: Ingresos menos Gastos reales comprobados
                ws_resumen.cell(row=row_num, column=11, value=f"=B{row_num}-E{row_num}")
        else:
            # Fórmula viva Arts. 100-110 / Arts. 206-215 LISR: Flujo de efectivo (Ingresos menos Gastos)
            ws_resumen.cell(row=row_num, column=11, value=f"=B{row_num}-E{row_num}")

        for col_idx in range(1, 12):
            c = ws_resumen.cell(row=row_num, column=col_idx)
            c.font = font_regular
            c.border = thin_border
            if col_idx == 1:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.number_format = num_format
                c.alignment = Alignment(horizontal="right", vertical="center")

            # Resaltado contable: IVA a cargo en rojo suave, saldo a favor en verde suave
            dif_iva_val = (iva16_em + iva8_em) - (iva16_rec + iva8_rec) - ret_iva_suf
            if col_idx == 9:
                if dif_iva_val > 0.01:
                    c.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                    c.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
                elif dif_iva_val < -0.01:
                    c.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                    c.font = Font(name="Calibri", size=10, bold=True, color="166534")
            elif col_idx == 11:
                c.fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
                c.font = Font(name="Calibri", size=10, bold=True, color="1E1B4B")
            elif curr_row % 2 == 0:
                c.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        curr_row += 1

    tot_label = f"TOTAL {MESES_NOMBRES[month].upper()}" if (month and month > 0) else "TOTAL ANUAL"
    ws_resumen.cell(row=tot_row, column=1, value=tot_label)
    ws_resumen.cell(row=tot_row, column=2, value=f"=SUM(B9:B{tot_row-1})")
    ws_resumen.cell(row=tot_row, column=3, value=f"=SUM(C9:C{tot_row-1})")
    ws_resumen.cell(row=tot_row, column=4, value=f"=SUM(D9:D{tot_row-1})")
    ws_resumen.cell(row=tot_row, column=5, value=f"=SUM(E9:E{tot_row-1})")
    ws_resumen.cell(row=tot_row, column=6, value=f"=SUM(F9:F{tot_row-1})")
    ws_resumen.cell(row=tot_row, column=7, value=f"=SUM(G9:G{tot_row-1})")
    ws_resumen.cell(row=tot_row, column=8, value=f"=SUM(H9:H{tot_row-1})")
    ws_resumen.cell(row=tot_row, column=9, value=f"=SUM(I9:I{tot_row-1})")
    ws_resumen.cell(row=tot_row, column=10, value=f"=SUM(J9:J{tot_row-1})")
    ws_resumen.cell(row=tot_row, column=11, value=f"=SUM(K9:K{tot_row-1})")

    for col_idx in range(1, 12):
        c = ws_resumen.cell(row=tot_row, column=col_idx)
        c.font = font_bold
        c.fill = fill_total
        c.border = total_border
        if col_idx == 1:
            c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.number_format = num_format
            c.alignment = Alignment(horizontal="right", vertical="center")

    # -------------------------------------------------------------
    # GRÁFICA NATIVA DE AUDITORÍA (IVA Trasladado vs Acreditable)
    # -------------------------------------------------------------
    try:
        from openpyxl.chart import BarChart, Reference
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = f"Comparativa de Flujo de IVA ({periodo_title})"
        chart.y_axis.title = "Importe ($ MXN)"
        chart.x_axis.title = "Mes"
        chart.width = 16
        chart.height = 11

        data_tras = Reference(ws_resumen, min_col=3, min_row=8, max_row=tot_row - 1)
        data_acred = Reference(ws_resumen, min_col=6, min_row=8, max_row=tot_row - 1)
        cats = Reference(ws_resumen, min_col=1, min_row=9, max_row=tot_row - 1)
        chart.add_data(data_tras, titles_from_data=True)
        chart.add_data(data_acred, titles_from_data=True)
        chart.set_categories(cats)
        ws_resumen.add_chart(chart, "M4")
    except Exception:
        pass

    # -------------------------------------------------------------
    # CÉDULA DE DICTAMEN FISCAL Y CONCLUSIONES PARA EL CONTADOR
    # -------------------------------------------------------------
    c_start = tot_row + 3
    ws_resumen.merge_cells(f"A{c_start}:K{c_start}")
    dictamen_title = ws_resumen.cell(row=c_start, column=1, value="CÉDULA DE AUDITORÍA FISCAL Y NOTAS DEL CONTADOR PÚBLICO")
    dictamen_title.font = font_header
    dictamen_title.fill = fill_navy
    dictamen_title.alignment = Alignment(horizontal="center", vertical="center")

    # Estatus IVA - Etiquetado dinámico
    r_iva = c_start + 1
    ws_resumen.merge_cells(f"A{r_iva}:F{r_iva}")
    ws_resumen.cell(row=r_iva, column=1, value="1. Estatus Definitivo de IVA del Periodo (Flujo de Efectivo):").font = font_bold
    ws_resumen.merge_cells(f"G{r_iva}:I{r_iva}")
    val_iva_status = ws_resumen.cell(row=r_iva, column=7, value=f'=IF(I{tot_row}>0.01, "IVA A CARGO (A ENTERAR AL SAT)", IF(I{tot_row}<-0.01, "SALDO A FAVOR (ACREDITABLE)", "SALDO DETERMINADO: $0.00"))')
    val_iva_status.font = font_bold
    val_iva_status.alignment = Alignment(horizontal="center")
    ws_resumen.merge_cells(f"J{r_iva}:K{r_iva}")
    c_iva_neto = ws_resumen.cell(row=r_iva, column=10, value=f"=ABS(I{tot_row})")
    c_iva_neto.font = font_bold
    c_iva_neto.number_format = num_format
    c_iva_neto.alignment = Alignment(horizontal="right")

    # ISR y Coeficiente / RESICO / Arrendamiento / RESICO PM
    r_isr = c_start + 2
    ws_resumen.merge_cells(f"A{r_isr}:F{r_isr}")
    if reg_code == '626':
        ws_resumen.cell(row=r_isr, column=1, value="2. Impuesto Sobre la Renta (RESICO PF - Art. 113-E LISR):").font = font_bold
        ws_resumen.merge_cells(f"G{r_isr}:I{r_isr}")
        val_isr_margen = ws_resumen.cell(row=r_isr, column=7, value='="Tasa RESICO: 1.0% a 2.5% Directa sobre Ingresos Cobrados"')
    elif reg_code == '606':
        arrend_sub = "Deducción Ciega del 35% sin comprobantes (Art. 115)" if opcion_arrendamiento == "ciega" else "Deducciones Comprobadas Reales con Facturas"
        ws_resumen.cell(row=r_isr, column=1, value=f"2. ISR Arrendamiento ({'Ciega 35%' if opcion_arrendamiento == 'ciega' else 'Gastos Reales'}):").font = font_bold
        ws_resumen.merge_cells(f"G{r_isr}:I{r_isr}")
        val_isr_margen = ws_resumen.cell(row=r_isr, column=7, value=f'="{arrend_sub}"')
    elif reg_code == 'resico_pm':
        ws_resumen.cell(row=r_isr, column=1, value="2. ISR RESICO Personas Morales (Flujo de Efectivo - Título VII Cap. XII):").font = font_bold
        ws_resumen.merge_cells(f"G{r_isr}:I{r_isr}")
        val_isr_margen = ws_resumen.cell(row=r_isr, column=7, value='="Flujo de Efectivo Acumulado: Cobrado menos Pagado"')
    else:
        ws_resumen.cell(row=r_isr, column=1, value="2. Base Gravable Estimada ISR (Act. Empresarial y Profesional):").font = font_bold
        ws_resumen.merge_cells(f"G{r_isr}:I{r_isr}")
        val_isr_margen = ws_resumen.cell(row=r_isr, column=7, value=f'="Margen Utilidad: " & TEXT(IF(B{tot_row}>0, K{tot_row}/B{tot_row}, 0), "0.0%")')
    val_isr_margen.font = font_bold
    val_isr_margen.alignment = Alignment(horizontal="center")
    ws_resumen.merge_cells(f"J{r_isr}:K{r_isr}")
    c_isr_util = ws_resumen.cell(row=r_isr, column=10, value=f"=K{tot_row}")
    c_isr_util.font = font_bold
    c_isr_util.number_format = num_format
    c_isr_util.alignment = Alignment(horizontal="right")

    # Retenciones Sufridas
    r_ret = c_start + 3
    ws_resumen.merge_cells(f"A{r_ret}:F{r_ret}")
    ws_resumen.cell(row=r_ret, column=1, value="3. Retenciones Sufridas por Clientes (Acreditamiento vs Impuesto):").font = font_bold
    ws_resumen.merge_cells(f"G{r_ret}:I{r_ret}")
    val_ret_desc = ws_resumen.cell(row=r_ret, column=7, value=f'="ISR: " & TEXT(J{tot_row}, "$#,##0.00") & " | IVA: " & TEXT(H{tot_row}, "$#,##0.00")')
    val_ret_desc.font = Font(name="Calibri", size=10, color="475569")
    val_ret_desc.alignment = Alignment(horizontal="center")
    ws_resumen.merge_cells(f"J{r_ret}:K{r_ret}")
    c_ret_tot = ws_resumen.cell(row=r_ret, column=10, value=f"=H{tot_row}+J{tot_row}")
    c_ret_tot.font = font_bold
    c_ret_tot.number_format = num_format
    c_ret_tot.alignment = Alignment(horizontal="right")

    # Fundamento Legal
    r_leg = c_start + 4
    ws_resumen.merge_cells(f"A{r_leg}:K{r_leg}")
    fund_reg = reg_info.get("fundamento", "LISR")
    fund_txt = f"Fundamento Legal: Determinación conforme a Arts. 1, 1-A, 5 LIVA y {fund_reg} ({reg_info.get('nombre', '')})"
    ws_resumen.cell(row=r_leg, column=1, value=fund_txt).font = Font(name="Calibri", size=9, italic=True, color="64748B")

    # Bordes en el bloque del dictamen
    for r_i in range(r_iva, r_leg + 1):
        for c_i in range(1, 12):
            cell_d = ws_resumen.cell(row=r_i, column=c_i)
            cell_d.border = thin_border
            if r_i in [r_iva, r_isr, r_ret] and c_i in [10, 11]:
                cell_d.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # Firmas
    r_sig1 = r_leg + 3
    r_sig2 = r_sig1 + 1
    r_sig3 = r_sig1 + 2
    ws_resumen.merge_cells(f"B{r_sig1}:E{r_sig1}")
    ws_resumen.cell(row=r_sig1, column=2, value="__________________________________________").alignment = Alignment(horizontal="center")
    ws_resumen.merge_cells(f"G{r_sig1}:J{r_sig1}")
    ws_resumen.cell(row=r_sig1, column=7, value="__________________________________________").alignment = Alignment(horizontal="center")

    ws_resumen.merge_cells(f"B{r_sig2}:E{r_sig2}")
    ws_resumen.cell(row=r_sig2, column=2, value="L.C. y C.P. Auditor / Contador Titular").font = font_bold
    ws_resumen.cell(row=r_sig2, column=2).alignment = Alignment(horizontal="center")
    ws_resumen.merge_cells(f"G{r_sig2}:J{r_sig2}")
    ws_resumen.cell(row=r_sig2, column=7, value="Contribuyente / Representante Legal").font = font_bold
    ws_resumen.cell(row=r_sig2, column=7).alignment = Alignment(horizontal="center")

    ws_resumen.merge_cells(f"B{r_sig3}:E{r_sig3}")
    ws_resumen.cell(row=r_sig3, column=2, value="Cédula Profesional / Firma Autógrafa").font = Font(name="Calibri", size=9, italic=True, color="64748B")
    ws_resumen.cell(row=r_sig3, column=2).alignment = Alignment(horizontal="center")
    ws_resumen.merge_cells(f"G{r_sig3}:J{r_sig3}")
    ws_resumen.cell(row=r_sig3, column=7, value="Aceptación y Entero de Cifras").font = Font(name="Calibri", size=9, italic=True, color="64748B")
    ws_resumen.cell(row=r_sig3, column=7).alignment = Alignment(horizontal="center")

    # ==========================================
    # PESTAÑA 2: Determinación Detallada de IVA
    # ==========================================
    ws_iva = wb.create_sheet(title="2_Determinacion_IVA")
    ws_iva.views.sheetView[0].showGridLines = True

    ws_iva["A1"] = f"DETERMINACIÓN DE IVA TRASLADADO VS ACREDITABLE - {periodo_title}"
    ws_iva["A1"].font = font_title

    headers_iva = [
        "Periodo", "Ventas Tasa 16%", "IVA Trasladado 16%", "Ventas Tasa 8%", "IVA Trasladado 8%",
        "Ventas Tasa 0%", "Ventas Exentas", "Total IVA Trasladado",
        "Compras Tasa 16%", "IVA Acreditable 16%", "Compras Tasa 8%", "IVA Acreditable 8%",
        "Total IVA Acreditable", "Retenciones IVA Sufridas", "IVA a Cargo / (Saldo a Favor)"
    ]

    ws_iva.append([])
    ws_iva.append(headers_iva)

    for col_idx in range(1, len(headers_iva) + 1):
        cell = ws_iva.cell(row=4, column=col_idx)
        cell.font = font_header
        cell.fill = fill_indigo
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    curr_iva_r = 5
    for m in meses_a_procesar:
        r = curr_iva_r
        em_mes = [x for x in datos_emitidas if x.get('mes') == m]
        rec_mes = [x for x in datos_recibidas if x.get('mes') == m]

        v_16 = sum(x['subtotal'] for x in em_mes if x['iva_16'] > 0)
        iva_16_t = sum(x['iva_16'] for x in em_mes)
        v_8 = sum(x['subtotal'] for x in em_mes if x['iva_8'] > 0)
        iva_8_t = sum(x['iva_8'] for x in em_mes)
        v_0 = sum(x['subtotal'] for x in em_mes if x['iva_0'] > 0)
        v_ex = sum(x['subtotal'] for x in em_mes if x['iva_exento'] > 0)

        # Gastos deducibles e IVA acreditable
        c_16 = sum(x.get('subtotal_deducible', x['subtotal']) for x in rec_mes if x.get('es_deducible', True) and x['iva_16'] > 0)
        iva_16_a = sum(x['iva_16'] for x in rec_mes if x.get('es_deducible', True))
        c_8 = sum(x.get('subtotal_deducible', x['subtotal']) for x in rec_mes if x.get('es_deducible', True) and x['iva_8'] > 0)
        iva_8_a = sum(x['iva_8'] for x in rec_mes if x.get('es_deducible', True))

        ret_iva = sum(x['ret_iva'] for x in em_mes)

        ws_iva.append([
            MESES_NOMBRES[m],
            v_16, iva_16_t, v_8, iva_8_t, v_0, v_ex,
            f"=C{r}+E{r}",
            c_16, iva_16_a, c_8, iva_8_a,
            f"=J{r}+L{r}",
            ret_iva,
            f"=H{r}-M{r}-N{r}"
        ])

        for col_idx in range(1, 16):
            c = ws_iva.cell(row=r, column=col_idx)
            c.font = font_regular
            c.border = thin_border
            if col_idx == 1:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.number_format = num_format
                c.alignment = Alignment(horizontal="right", vertical="center")

            # Colores contables para identificar valores clave
            if col_idx == 8:  # Total IVA Trasladado
                c.fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
                c.font = Font(name="Calibri", size=10, bold=True, color="312E81")
            elif col_idx == 13:  # Total IVA Acreditable
                c.fill = PatternFill(start_color="F0FDFA", end_color="F0FDFA", fill_type="solid")
                c.font = Font(name="Calibri", size=10, bold=True, color="134E4A")
            elif col_idx == 14:  # Retenciones IVA Sufridas
                c.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                c.font = Font(name="Calibri", size=10, bold=True, color="78350F")
            elif col_idx == 15:  # IVA a Cargo / Saldo a Favor
                net_iva_val = (iva_16_t + iva_8_t) - (iva_16_a + iva_8_a) - ret_iva
                if net_iva_val > 0.01:
                    c.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                    c.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
                elif net_iva_val < -0.01:
                    c.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                    c.font = Font(name="Calibri", size=10, bold=True, color="166534")
            elif curr_iva_r % 2 == 0:
                c.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        curr_iva_r += 1

    tot_iva_row = curr_iva_r
    tot_iva_label = f"TOTAL {MESES_NOMBRES[month].upper()}" if (month and month > 0) else "TOTAL ANUAL"
    ws_iva.append([
        tot_iva_label,
        f"=SUM(B5:B{tot_iva_row-1})", f"=SUM(C5:C{tot_iva_row-1})", f"=SUM(D5:D{tot_iva_row-1})", f"=SUM(E5:E{tot_iva_row-1})",
        f"=SUM(F5:F{tot_iva_row-1})", f"=SUM(G5:G{tot_iva_row-1})", f"=SUM(H5:H{tot_iva_row-1})", f"=SUM(I5:I{tot_iva_row-1})",
        f"=SUM(J5:J{tot_iva_row-1})", f"=SUM(K5:K{tot_iva_row-1})", f"=SUM(L5:L{tot_iva_row-1})", f"=SUM(M5:M{tot_iva_row-1})",
        f"=SUM(N5:N{tot_iva_row-1})", f"=SUM(O5:O{tot_iva_row-1})"
    ])
    for col_idx in range(1, 16):
        c = ws_iva.cell(row=tot_iva_row, column=col_idx)
        c.font = font_bold
        c.fill = fill_total
        c.border = total_border
        if col_idx == 1:
            c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.number_format = num_format
            c.alignment = Alignment(horizontal="right", vertical="center")

    # ==========================================
    # PESTAÑA 3: Concentrado de Retenciones (ISR/IVA)
    # ==========================================
    ws_ret = wb.create_sheet(title="3_Concentrado_Retenciones")
    ws_ret.views.sheetView[0].showGridLines = True

    ws_ret["A1"] = f"CONCENTRADO DE RETENCIONES DE ISR E IVA POR TERCEROS - {periodo_title}"
    ws_ret["A1"].font = font_title

    headers_ret = [
        "Tipo Operación", "RFC Tercero", "Nombre / Razón Social", "Subtotal Acumulado",
        "Retención ISR Acumulada", "Retención IVA Acumulada", "Total Operaciones"
    ]
    ws_ret.append([])
    ws_ret.append(headers_ret)

    for col_idx in range(1, len(headers_ret) + 1):
        cell = ws_ret.cell(row=4, column=col_idx)
        cell.font = font_header
        cell.fill = fill_amber
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    agrupado = {}
    datos_em_ret = [d for d in datos_emitidas if d.get('mes') == month] if (month and month > 0) else datos_emitidas
    datos_rec_ret = [d for d in datos_recibidas if d.get('mes') == month] if (month and month > 0) else datos_recibidas

    for d in datos_em_ret:
        if d['ret_isr'] > 0 or d['ret_iva'] > 0:
            key = ('Emitida (Cliente)', d['receptor_rfc'], d['receptor_nombre'])
            if key not in agrupado:
                agrupado[key] = {'subtotal': 0.0, 'ret_isr': 0.0, 'ret_iva': 0.0, 'count': 0}
            agrupado[key]['subtotal'] += d['subtotal']
            agrupado[key]['ret_isr'] += d['ret_isr']
            agrupado[key]['ret_iva'] += d['ret_iva']
            agrupado[key]['count'] += 1

    for d in datos_rec_ret:
        if d['ret_isr'] > 0 or d['ret_iva'] > 0:
            key = ('Recibida (Proveedor)', d['emisor_rfc'], d['emisor_nombre'])
            if key not in agrupado:
                agrupado[key] = {'subtotal': 0.0, 'ret_isr': 0.0, 'ret_iva': 0.0, 'count': 0}
            agrupado[key]['subtotal'] += d['subtotal']
            agrupado[key]['ret_isr'] += d['ret_isr']
            agrupado[key]['ret_iva'] += d['ret_iva']
            agrupado[key]['count'] += 1

    ret_row = 5
    if not agrupado:
        ws_ret.append(["Sin registros de retención de ISR o IVA en el periodo", "", "", 0.0, 0.0, 0.0, 0])
        ret_row = 6
    else:
        for (tipo_op, rfc_tercero, nombre_tercero), vals in agrupado.items():
            ws_ret.append([
                tipo_op, rfc_tercero, nombre_tercero,
                vals['subtotal'], vals['ret_isr'], vals['ret_iva'], vals['count']
            ])
            for col_idx in range(1, 8):
                c = ws_ret.cell(row=ret_row, column=col_idx)
                c.font = font_regular
                c.border = thin_border
                if col_idx in [4, 5, 6]:
                    c.number_format = num_format
                    c.alignment = Alignment(horizontal="right", vertical="center")
                    if col_idx in [5, 6] and (vals['ret_isr'] > 0 or vals['ret_iva'] > 0):
                        c.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                        c.font = Font(name="Calibri", size=10, bold=True, color="78350F")
                elif col_idx in [1, 2, 7]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    if col_idx == 2:
                        c.font = font_bold
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")
            ret_row += 1

    ws_ret.append([
        "TOTALES", "", "", f"=SUM(D5:D{ret_row-1})", f"=SUM(E5:E{ret_row-1})", f"=SUM(F5:F{ret_row-1})", f"=SUM(G5:G{ret_row-1})"
    ])
    for col_idx in range(1, 8):
        c = ws_ret.cell(row=ret_row, column=col_idx)
        c.font = font_bold
        c.fill = fill_total
        c.border = total_border
        if col_idx in [4, 5, 6]:
            c.number_format = num_format
            c.alignment = Alignment(horizontal="right", vertical="center")
        elif col_idx in [1, 2, 7]:
            c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.alignment = Alignment(horizontal="left", vertical="center")

    # ==========================================
    # PESTAÑA 4: Auxiliar Detalle Emitidas
    # ==========================================
    ws_em = wb.create_sheet(title="4_Detalle_Ingresos")
    ws_em.views.sheetView[0].showGridLines = True
    headers_aux = ["Folio Fiscal (UUID)", "Fecha", "Mes", "RFC Receptor", "Nombre Receptor", "Método Pago", "Subtotal", "IVA 16%", "IVA 8%", "Ret ISR", "Ret IVA", "Total", "Auditoría LISR"]
    ws_em.append(headers_aux)
    for col_idx in range(1, len(headers_aux) + 1):
        cell = ws_em.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_indigo
        cell.alignment = Alignment(horizontal="center", vertical="center")

    em_idx = 2
    datos_em_det = [d for d in datos_emitidas if d.get('mes') == month] if (month and month > 0) else datos_emitidas
    for d in datos_em_det:
        ws_em.append([
            d['uuid'], d['fecha'], MESES_NOMBRES[d['mes']], d['receptor_rfc'], d['receptor_nombre'],
            d['metodo_pago'], d['subtotal'], d['iva_16'], d['iva_8'], d['ret_isr'], d['ret_iva'],
            f"=G{em_idx}+H{em_idx}+I{em_idx}-J{em_idx}-K{em_idx}",
            d.get('riesgo_deducibilidad', '')
        ])
        for col_idx in range(1, 14):
            c = ws_em.cell(row=em_idx, column=col_idx)
            c.font = font_regular
            c.border = thin_border
            if col_idx in [1, 2, 3, 4, 6]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 5:
                c.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx >= 7 and col_idx <= 12:
                c.number_format = num_format
                c.alignment = Alignment(horizontal="right", vertical="center")
                if col_idx == 12:
                    c.font = font_bold
            if col_idx == 13 and d.get('riesgo_deducibilidad'):
                c.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                c.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
            elif em_idx % 2 == 0 and not (col_idx == 13 and d.get('riesgo_deducibilidad')):
                c.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        em_idx += 1

    if em_idx > 2:
        ws_em.append([
            "TOTAL GENERAL", "", "", "", "", "",
            f"=SUM(G2:G{em_idx-1})", f"=SUM(H2:H{em_idx-1})", f"=SUM(I2:I{em_idx-1})",
            f"=SUM(J2:J{em_idx-1})", f"=SUM(K2:K{em_idx-1})", f"=SUM(L2:L{em_idx-1})",
            f"{len(datos_em_det)} Facturas"
        ])
        for col_idx in range(1, 14):
            c = ws_em.cell(row=em_idx, column=col_idx)
            c.font = font_bold
            c.fill = fill_total
            c.border = total_border
            if col_idx in [1, 2, 3, 4, 6]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx >= 7 and col_idx <= 12:
                c.number_format = num_format
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")

    # ==========================================
    # PESTAÑA 5: Auxiliar Detalle Recibidas
    # ==========================================
    ws_rec = wb.create_sheet(title="5_Detalle_Gastos")
    ws_rec.views.sheetView[0].showGridLines = True
    headers_aux_rec = ["Folio Fiscal (UUID)", "Fecha", "Mes", "RFC Emisor", "Nombre Emisor", "Método Pago", "Subtotal", "IVA 16%", "IVA 8%", "Ret ISR", "Ret IVA", "Total", "Auditoría LISR"]
    ws_rec.append(headers_aux_rec)
    for col_idx in range(1, len(headers_aux_rec) + 1):
        cell = ws_rec.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_teal
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rec_idx = 2
    datos_rec_det = [d for d in datos_recibidas if d.get('mes') == month] if (month and month > 0) else datos_recibidas
    for d in datos_rec_det:
        ws_rec.append([
            d['uuid'], d['fecha'], MESES_NOMBRES[d['mes']], d['emisor_rfc'], d['emisor_nombre'],
            d['metodo_pago'], d['subtotal'], d['iva_16'], d['iva_8'], d['ret_isr'], d['ret_iva'],
            f"=G{rec_idx}+H{rec_idx}+I{rec_idx}-J{rec_idx}-K{rec_idx}",
            d.get('riesgo_deducibilidad', '')
        ])
        for col_idx in range(1, 14):
            c = ws_rec.cell(row=rec_idx, column=col_idx)
            c.font = font_regular
            c.border = thin_border
            if col_idx in [1, 2, 3, 4, 6]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 5:
                c.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx >= 7 and col_idx <= 12:
                c.number_format = num_format
                c.alignment = Alignment(horizontal="right", vertical="center")
                if col_idx == 12:
                    c.font = font_bold
            if col_idx == 13 and d.get('riesgo_deducibilidad'):
                c.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                c.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
            elif rec_idx % 2 == 0 and not (col_idx == 13 and d.get('riesgo_deducibilidad')):
                c.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        rec_idx += 1

    if rec_idx > 2:
        ws_rec.append([
            "TOTAL GENERAL", "", "", "", "", "",
            f"=SUM(G2:G{rec_idx-1})", f"=SUM(H2:H{rec_idx-1})", f"=SUM(I2:I{rec_idx-1})",
            f"=SUM(J2:J{rec_idx-1})", f"=SUM(K2:K{rec_idx-1})", f"=SUM(L2:L{rec_idx-1})",
            f"{len(datos_rec_det)} Facturas"
        ])
        for col_idx in range(1, 14):
            c = ws_rec.cell(row=rec_idx, column=col_idx)
            c.font = font_bold
            c.fill = fill_total
            c.border = total_border
            if col_idx in [1, 2, 3, 4, 6]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx >= 7 and col_idx <= 12:
                c.number_format = num_format
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")

    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    if report_type == 'retenciones':
        wb.active = ws_ret
    elif report_type == 'iva_mensual':
        wb.active = ws_iva
    else:
        wb.active = ws_resumen

    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if not val.startswith("="):
                    max_len = max(max_len, len(val))
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path

def convertir_registro_sat(r, tipo_op, default_year):
    """Normaliza un registro extraído del SAT al formato interno del análisis"""
    if not isinstance(r, dict):
        return None
    estado = str(r.get('Estado del Comprobante') or r.get('estado') or 'Vigente').strip()
    if 'cancelado' in estado.lower():
        return None  # Excluir facturas canceladas en el SAT
    
    uuid = str(r.get('Folio') or r.get('uuid') or '').strip()
    fecha_raw = str(r.get('Fecha de Emisión') or r.get('fecha') or '').strip()
    
    # Extraer mes y año
    mes = int(r.get('NumMes') or r.get('mes') or 0)
    anio = int(r.get('anio') or default_year)
    if not mes and fecha_raw:
        try:
            if 'T' in fecha_raw:
                dt = datetime.fromisoformat(fecha_raw.split('T')[0])
                mes = dt.month
                anio = dt.year
            elif '-' in fecha_raw:
                parts = fecha_raw.split('-')
                if len(parts) == 3:
                    if len(parts[0]) == 4:
                        anio = int(parts[0])
                        mes = int(parts[1])
                    else:
                        anio = int(parts[2])
                        for m_idx, m_name in enumerate(MESES_NOMBRES):
                            if m_name and m_name.lower() in parts[1].lower():
                                mes = m_idx
                                break
        except Exception:
            pass

    nombre = str(r.get('Nombre o Razón Social') or r.get('nombre') or '')
    rfc = str(r.get('RFC') or r.get('rfc') or '')
    total = float(r.get('Total') or r.get('total') or 0.0)
    subtotal = float(r.get('Importe') or r.get('subtotal') or 0.0)
    efecto = str(r.get('Efecto del Comprobante') or r.get('tipo') or 'I')
    forma_pago_raw = r.get('FormaPago') or r.get('forma_pago') or ''
    metodo_pago_raw = r.get('MetodoPago') or r.get('metodo_pago') or ''

    # Normalización para detección fiscal rigurosa
    import unicodedata
    def _norm(s):
        return unicodedata.normalize('NFKD', str(s or '')).encode('ASCII', 'ignore').decode('utf-8').lower().strip()

    c_efecto = _norm(efecto)
    c_nombre = _norm(nombre)

    is_nomina = 'nomina' in c_efecto or c_efecto == 'n' or 'nmina' in c_efecto
    is_educacion = any(w in c_nombre for w in ['instituto', 'colegio', 'escuela', 'universidad', 'educat', 'educac', 'ensenan', 'liceo', 'kinder', 'primaria', 'secundaria', 'preparatoria', 'academ', 'formacion', 'docen', 'pedagog']) and not ('seguro social' in c_nombre or 'imss' in c_nombre or 'issste' in c_nombre)
    is_medico = any(w in c_nombre for w in ['hospital', 'clinica', 'medico', 'sanatorio', 'salud', 'laboratorio', 'diagnostico'])
    is_pago = c_efecto == 'p' or 'pago' in c_efecto
    is_egreso = c_efecto == 'e' or 'egreso' in c_efecto
    is_gasolina = any(k in c_nombre for k in ['gasolin', 'combustib', 'hidrocarb', 'diesel', 'estacion de serv', 'oxxo gas', 'hidrosina', 'petro', 'gas popo', 'origas', 'gastoga']) or '151015' in str(r)

    if is_pago:
        metodo_pago = 'REP'
        forma_pago = forma_pago_raw or 'REP'
    elif is_egreso:
        metodo_pago = 'Egreso'
        forma_pago = forma_pago_raw or 'EGRESO'
    elif is_nomina:
        metodo_pago = 'Nómina'
        forma_pago = forma_pago_raw or '99'
    else:
        metodo_pago = str(metodo_pago_raw or 'PUE').strip()
        forma_pago = str(forma_pago_raw or 'SIN_XML').strip()

    if forma_pago and forma_pago.isdigit():
        forma_pago = forma_pago.zfill(2)

    iva_exento = 0.0
    riesgo_deducibilidad = ""

    if is_nomina:
        subtotal = total
        iva_16 = 0.0
        iva_8 = 0.0
        riesgo_deducibilidad = "ℹ️ï¸ CFDI de Nómina (Sueldos y Salarios - No Causa IVA)"
    elif is_educacion:
        subtotal = total
        iva_16 = 0.0
        iva_8 = 0.0
        iva_exento = total
        riesgo_deducibilidad = "ℹ️ï¸ Servicio de Enseñanza Exento (Art. 15 Frac. IV LIVA)"
    elif is_medico:
        subtotal = total
        iva_16 = 0.0
        iva_8 = 0.0
        iva_exento = total
        riesgo_deducibilidad = "ℹ️ï¸ Honorarios Médicos Exentos (Art. 15 Frac. XIV LIVA)"
    elif is_pago:
        subtotal = 0.0
        iva_16 = 0.0
        iva_8 = 0.0
        total = 0.0
        riesgo_deducibilidad = "ℹ️ï¸ Complemento para Recepción de Pagos (REP)"
    elif is_egreso:
        if subtotal == 0.0 and total > 0.0:
            subtotal = round(total / 1.16, 2)
        iva_16 = float(r.get('IVA al 16%') or r.get('iva_16') or 0.0)
        if iva_16 == 0.0 and subtotal > 0.0 and total > subtotal:
            iva_16 = round(total - subtotal, 2)
        subtotal = -abs(subtotal)
        iva_16 = -abs(iva_16)
        total = -abs(total)
        riesgo_deducibilidad = "ℹ️ï¸ Nota de Crédito de Proveedor (Disminuye Gasto - Art. 25 LISR)"
    else:
        if subtotal == 0.0 and total > 0.0:
            subtotal = round(total / 1.16, 2)
        iva_16 = float(r.get('IVA al 16%') or r.get('iva_16') or 0.0)
        # Si venía de JSON previo con IVA artificial en Nómina o Escuelas, corregirlo
        if iva_16 == 0.0 and subtotal > 0.0 and total > subtotal:
            iva_16 = round(total - subtotal, 2)

    iva_8 = float(r.get('iva_8') or 0.0)
    ret_isr = float(r.get('ret_isr') or 0.0)
    ret_iva = float(r.get('ret_iva') or 0.0)

    # Evaluación fiscal y listas negras SAT (EFOS Art. 69-B CFF & Reglas LISR)
    efos_info = check_rfc_efos(rfc) if tipo_op == 'recibidas' else None
    is_efos_definitivo = bool(efos_info and efos_info.get('es_definitivo'))

    es_deducible = True
    no_deducible_tipo = ""

    if tipo_op == 'recibidas':
        if is_efos_definitivo:
            es_deducible = False
            no_deducible_tipo = "efos_definitivo"
            riesgo_deducibilidad = f"🚨 EFOS Definitivo (Art. 69-B CFF): Factura No Deducible"
            iva_16 = 0.0
            iva_8 = 0.0
        elif is_pago:
            es_deducible = False
            no_deducible_tipo = "rep_control"
        elif is_egreso:
            es_deducible = True
        elif is_gasolina and forma_pago == '01':
            es_deducible = False
            no_deducible_tipo = "gasolina_efectivo"
            riesgo_deducibilidad = "⚠️ï¸ Gasolina en Efectivo: 100% No Deducible (Art. 27 Fracc. III LISR)"
            iva_16 = 0.0
            iva_8 = 0.0
        elif forma_pago == '01' and total > 2000.0:
            es_deducible = False
            no_deducible_tipo = "efectivo_mayor_2k"
            riesgo_deducibilidad = "⚠️ï¸ Pago en Efectivo > $2,000: No Deducible (Art. 27 Fracc. III LISR)"
            iva_16 = 0.0
            iva_8 = 0.0
        elif is_nomina:
            es_deducible = True
            riesgo_deducibilidad = "ℹ️ï¸ CFDI de Nómina (Sueldos y Salarios - IVA $0.00)"
        elif is_educacion:
            es_deducible = True
            riesgo_deducibilidad = "ℹ️ï¸ Servicio de Enseñanza Exento (Art. 15 Frac. IV LIVA)"
        elif is_medico:
            es_deducible = True
            riesgo_deducibilidad = "ℹ️ï¸ Honorarios Médicos Exentos (Art. 15 Frac. XIV LIVA)"

    subtotal_deducible = subtotal if es_deducible else 0.0

    return {
        'uuid': uuid,
        'fecha': fecha_raw,
        'mes': mes,
        'anio': anio,
        'tipo': efecto[0] if efecto else 'I',
        'metodo_pago': metodo_pago,
        'forma_pago': forma_pago,
        'emisor_rfc': 'CONTRIBUYENTE' if tipo_op == 'emitidas' else rfc,
        'emisor_nombre': 'CONTRIBUYENTE' if tipo_op == 'emitidas' else nombre,
        'receptor_rfc': rfc if tipo_op == 'emitidas' else 'CONTRIBUYENTE',
        'receptor_nombre': nombre if tipo_op == 'emitidas' else 'CONTRIBUYENTE',
        'subtotal': subtotal,
        'subtotal_deducible': subtotal_deducible,
        'descuento': 0.0,
        'iva_16': iva_16,
        'iva_8': iva_8,
        'iva_0': 0.0,
        'iva_exento': iva_exento,
        'ret_isr': ret_isr,
        'ret_iva': ret_iva,
        'total': total,
        'es_deducible': es_deducible,
        'no_deducible_tipo': no_deducible_tipo,
        'is_gasolina': is_gasolina,
        'efos_info': efos_info,
        'riesgo_deducibilidad': riesgo_deducibilidad
    }

def main():
    parser = argparse.ArgumentParser(description="Generador del Papel de Trabajo de Análisis de IVA e ISR")
    parser.add_argument("--year", type=int, default=2026, help="Ejercicio fiscal a analizar")
    parser.add_argument("--month", type=int, default=0, help="Mes específico (1-12) o 0 para anual")
    parser.add_argument("--zone", type=str, default="general", choices=["general", "fronteriza_norte", "fronteriza_sur"], help="Zona fiscal")
    parser.add_argument("--type", type=str, default="completo", help="Tipo de reporte")
    parser.add_argument("--regime", type=str, default="612", help="Régimen fiscal del contribuyente (código SAT de 3 dígitos o alias)")
    parser.add_argument("--arrendamiento-opcion", type=str, default="ciega", choices=["ciega", "real"], help="Opción para Arrendamiento 606: ciega (35%%) o real (gastos)")
    parser.add_argument("--xml-dir", type=str, default=None, help="Directorio conteniendo archivos XML o paquetes ZIP")
    parser.add_argument("--invoices-dir", type=str, default=None, help="Directorio conteniendo archivos JSON sincronizados del SAT")
    parser.add_argument("--emitidas-json", type=str, default=None, help="Ruta directa a facturas emitidas en JSON")
    parser.add_argument("--recibidas-json", type=str, default=None, help="Ruta directa a facturas recibidas en JSON")
    parser.add_argument("--rfc", type=str, default=None, help="RFC del contribuyente activo para clasificar facturas emitidas y recibidas")
    parser.add_argument("--metrics-json", type=str, default=None, help="Ruta de archivo para guardar las métricas directamente en disco")
    parser.add_argument("--output", type=str, required=True, help="Ruta del archivo Excel de salida")
    args = parser.parse_args()

    reg_code = normalizar_regimen_code(args.regime)
    reg_info = get_regimen_info(reg_code)

    emit_event({"type": "progress", "percent": 15, "msg": f"Iniciando motor de Análisis Fiscal ({reg_info.get('short_label', reg_code)})..."})
    time.sleep(0.3)

    # 1. Buscar en directorios con archivos XML y ZIPs (Fuente Primaria Oficial y Verificada)
    search_dirs = []
    if args.xml_dir and os.path.exists(args.xml_dir):
        search_dirs.append(args.xml_dir)
    if args.invoices_dir and os.path.exists(args.invoices_dir) and args.invoices_dir not in search_dirs:
        search_dirs.append(args.invoices_dir)

    xml_emitidas = {}
    xml_recibidas = {}
    all_xmls_by_uuid = {}
    rfc_target = (args.rfc or '').upper().strip()

    import zipfile
    for target_dir in search_dirs:
        emit_event({"type": "progress", "percent": 30, "msg": f"Analizando comprobantes XML oficiales en {os.path.basename(target_dir)}..."})
        for root_dir, dirs, files in os.walk(target_dir):
            for file in files:
                file_path = os.path.join(root_dir, file)
                if file.lower().endswith(".xml"):
                    try:
                        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                            content = f.read()
                        parsed = parse_cfdi_xml(content)
                        if parsed and parsed.get('uuid'):
                            u = parsed['uuid'].upper().strip()
                            all_xmls_by_uuid[u] = parsed
                            for rel_u in parsed.get('relacionados', []):
                                if rel_u and rel_u not in all_xmls_by_uuid:
                                    rel_parsed = dict(parsed)
                                    rel_parsed['uuid'] = rel_u
                                    rel_parsed['uuid_sustituto'] = u
                                    rel_parsed['is_sustitucion'] = True
                                    all_xmls_by_uuid[rel_u] = rel_parsed
                            if parsed.get('anio') == args.year:
                                if rfc_target:
                                    if parsed.get('emisor_rfc', '').upper() == rfc_target:
                                        xml_emitidas[u] = parsed
                                    elif parsed.get('receptor_rfc', '').upper() == rfc_target:
                                        xml_recibidas[u] = parsed
                                    elif "emitidas" in file_path.lower() or "ingreso" in file_path.lower():
                                        xml_emitidas[u] = parsed
                                    else:
                                        xml_recibidas[u] = parsed
                                elif "emitidas" in file_path.lower() or "ingreso" in file_path.lower():
                                    xml_emitidas[u] = parsed
                                else:
                                    xml_recibidas[u] = parsed
                    except Exception:
                        pass
                elif file.lower().endswith(".zip"):
                    try:
                        with zipfile.ZipFile(file_path, 'r') as z:
                            for zname in z.namelist():
                                if zname.lower().endswith(".xml"):
                                    with z.open(zname) as zf:
                                        content = zf.read().decode('utf-8', errors='ignore')
                                        parsed = parse_cfdi_xml(content)
                                        if parsed and parsed.get('uuid'):
                                            u = parsed['uuid'].upper().strip()
                                            all_xmls_by_uuid[u] = parsed
                                            for rel_u in parsed.get('relacionados', []):
                                                if rel_u and rel_u not in all_xmls_by_uuid:
                                                    rel_parsed = dict(parsed)
                                                    rel_parsed['uuid'] = rel_u
                                                    rel_parsed['uuid_sustituto'] = u
                                                    rel_parsed['is_sustitucion'] = True
                                                    all_xmls_by_uuid[rel_u] = rel_parsed
                                            if parsed.get('anio') == args.year:
                                                if rfc_target:
                                                    if parsed.get('emisor_rfc', '').upper() == rfc_target:
                                                        xml_emitidas[u] = parsed
                                                    elif parsed.get('receptor_rfc', '').upper() == rfc_target:
                                                        xml_recibidas[u] = parsed
                                                    elif "emitidas" in zname.lower() or "ingreso" in zname.lower():
                                                        xml_emitidas[u] = parsed
                                                    else:
                                                        xml_recibidas[u] = parsed
                                                elif "emitidas" in zname.lower() or "ingreso" in zname.lower():
                                                    xml_emitidas[u] = parsed
                                                else:
                                                    xml_recibidas[u] = parsed
                    except Exception:
                        pass

    # 2. Leer facturas desde archivos JSON sincronizados del SAT (Deduplicando estrictamente por UUID)
    json_emitidas_paths = []
    json_recibidas_paths = []

    if args.emitidas_json and os.path.exists(args.emitidas_json):
        json_emitidas_paths.append(args.emitidas_json)
    if args.recibidas_json and os.path.exists(args.recibidas_json):
        json_recibidas_paths.append(args.recibidas_json)

    if args.invoices_dir and os.path.exists(args.invoices_dir):
        p_em = os.path.join(args.invoices_dir, f"{args.year}_emitidas.json")
        p_rec = os.path.join(args.invoices_dir, f"{args.year}_recibidas.json")
        if os.path.exists(p_em) and p_em not in json_emitidas_paths:
            json_emitidas_paths.append(p_em)
        if os.path.exists(p_rec) and p_rec not in json_recibidas_paths:
            json_recibidas_paths.append(p_rec)

    for p in json_emitidas_paths:
        try:
            with open(p, 'r', encoding='utf-8') as f:
                raw = json.load(f)
                reg_list = raw if isinstance(raw, list) else raw.get('registros', [])
                for item in reg_list:
                    u = (item.get('Folio') or item.get('uuid') or '').upper().strip()
                    if u and u in xml_emitidas:
                        continue  # Ya contamos con el CFDI XML original con datos fiscales fidedignos
                    if u and u in all_xmls_by_uuid:
                        parsed = all_xmls_by_uuid[u]
                        if parsed.get('is_sustitucion'):
                            conv = convertir_registro_sat(item, 'emitidas', args.year)
                            if conv:
                                conv['forma_pago'] = parsed.get('forma_pago', '04')
                                conv['metodo_pago'] = parsed.get('metodo_pago', 'PUE')
                                conv['has_xml'] = True
                                xml_emitidas[u] = conv
                                continue
                        elif parsed.get('anio') == args.year or parsed.get('anio') == 0:
                            xml_emitidas[u] = parsed
                            continue
                    conv = convertir_registro_sat(item, 'emitidas', args.year)
                    if conv and (conv['anio'] == args.year or conv['anio'] == 0):
                        if u not in xml_emitidas:
                            xml_emitidas[u] = conv
        except Exception:
            pass

    for p in json_recibidas_paths:
        try:
            with open(p, 'r', encoding='utf-8') as f:
                raw = json.load(f)
                reg_list = raw if isinstance(raw, list) else raw.get('registros', [])
                for item in reg_list:
                    u = (item.get('Folio') or item.get('uuid') or '').upper().strip()
                    if u and u in xml_recibidas:
                        continue  # Ya contamos con el CFDI XML original con datos fiscales fidedignos
                    if u and u in all_xmls_by_uuid:
                        parsed = all_xmls_by_uuid[u]
                        if parsed.get('is_sustitucion'):
                            conv = convertir_registro_sat(item, 'recibidas', args.year)
                            if conv:
                                conv['forma_pago'] = parsed.get('forma_pago', '04')
                                conv['metodo_pago'] = parsed.get('metodo_pago', 'PUE')
                                conv['has_xml'] = True
                                conv['es_deducible'] = True
                                if parsed.get('is_gasolina'):
                                    conv['is_gasolina'] = True
                                xml_recibidas[u] = conv
                                continue
                        elif parsed.get('anio') == args.year or parsed.get('anio') == 0:
                            xml_recibidas[u] = parsed
                            continue
                    conv = convertir_registro_sat(item, 'recibidas', args.year)
                    if conv and (conv['anio'] == args.year or conv['anio'] == 0):
                        if u not in xml_recibidas:
                            xml_recibidas[u] = conv
        except Exception:
            pass

    datos_emitidas = list(xml_emitidas.values())
    datos_recibidas = list(xml_recibidas.values())

    # Ordenar cronológicamente de forma estricta
    def get_item_sort_key(d):
        raw = str(d.get('fecha_raw') or '')
        if raw and 'T' in raw:
            return raw
        f = str(d.get('fecha') or '')
        if 'T' in f:
            return f
        if '-' in f:
            parts = f.split('-')
            if len(parts) == 3:
                if len(parts[0]) == 4:
                    return f"{parts[0]}-{parts[1].zfill(2)}-{parts[2].zfill(2)}"
                else:
                    day = parts[0].zfill(2)
                    mes = str(d.get('mes', 1)).zfill(2)
                    yr = parts[2]
                    return f"{yr}-{mes}-{day}"
        return f"{d.get('anio', 2025)}-{str(d.get('mes', 1)).zfill(2)}-01"

    datos_emitidas.sort(key=get_item_sort_key)
    datos_recibidas.sort(key=get_item_sort_key)

    # Integración de Nómina: Trasladar comprobantes emitidos de tipo "N" a deducciones de gasto
    reales_emitidas = []
    for d in datos_emitidas:
        if d.get('tipo') == 'N' or 'nomina' in str(d.get('riesgo_deducibilidad', '')).lower():
            d['es_nomina_emitida'] = True
            d['subtotal_deducible'] = d['total'] if d['total'] > 0 else d['subtotal']
            d['subtotal'] = d['subtotal_deducible']
            d['iva_16'] = 0.0
            d['iva_8'] = 0.0
            d['iva_exento'] = 0.0
            d['es_deducible'] = True
            d['categoria'] = 'nomina'
            d['riesgo_deducibilidad'] = "ℹ️ï¸ CFDI de Nómina Emitida (Deducción Autorizada Sueldos y Salarios - IVA $0.00)"
            datos_recibidas.append(d)
        else:
            reales_emitidas.append(d)
    datos_emitidas = reales_emitidas

    # Notificación de estado de datos
    has_sync_files = bool(args.invoices_dir and os.path.exists(args.invoices_dir) and (
        os.path.exists(os.path.join(args.invoices_dir, f"{args.year}_emitidas.json")) or
        os.path.exists(os.path.join(args.invoices_dir, f"{args.year}_recibidas.json")) or
        os.path.exists(os.path.join(args.invoices_dir, f"{args.year}_meta.json"))
    ))
    has_real_data = len(datos_emitidas) > 0 or len(datos_recibidas) > 0 or has_sync_files
    if has_real_data:
        emit_event({
            "type": "progress",
            "percent": 50,
            "msg": f"Comprobantes auditados: {len(datos_emitidas)} ingresos y {len(datos_recibidas)} gastos/nóminas (Deduplicación rigurosa activa)."
        })
    else:
        emit_event({
            "type": "progress",
            "percent": 50,
            "msg": f"Sin comprobantes oficiales registrados para el ejercicio {args.year}. Se presentará cédula en ceros."
        })

    def procesar_flujo(raw_list):
        pue = []
        ppd = {}
        reps = []
        for d in raw_list:
            if d.get('tipo') == 'P':
                reps.append(d)
            elif d.get('metodo_pago') == 'PPD':
                ppd[d.get('uuid', '').upper().strip()] = d
            else:
                pue.append(d)
        
        flujo = list(pue)
        for rep in reps:
            for dr in rep.get('doctos_relacionados', []):
                dr_u = dr.get('uuid', '').upper().strip()
                parent = ppd.get(dr_u)
                if parent:
                    synth = parent.copy()
                    synth['fecha'] = rep.get('fecha', parent.get('fecha'))
                    synth['mes'] = rep.get('mes', parent.get('mes'))
                    synth['anio'] = rep.get('anio', parent.get('anio'))
                    synth['forma_pago'] = rep.get('forma_pago', parent.get('forma_pago', '03'))
                    
                    if dr['iva_16'] == 0 and dr['iva_8'] == 0 and parent['total'] > 0:
                        prop = dr['imp_pagado'] / parent['total']
                        synth['subtotal'] = parent['subtotal'] * prop
                        synth['subtotal_deducible'] = parent.get('subtotal_deducible', parent['subtotal']) * prop
                        synth['iva_16'] = parent['iva_16'] * prop
                        synth['iva_8'] = parent['iva_8'] * prop
                        synth['ret_isr'] = parent['ret_isr'] * prop
                        synth['ret_iva'] = parent['ret_iva'] * prop
                    else:
                        synth['subtotal'] = round(dr['imp_pagado'] / 1.16, 2) if dr['iva_16'] > 0 else dr['imp_pagado']
                        synth['subtotal_deducible'] = synth['subtotal'] if parent.get('es_deducible', True) else 0.0
                        synth['iva_16'] = dr['iva_16']
                        synth['iva_8'] = dr['iva_8']
                        synth['ret_isr'] = 0.0
                        synth['ret_iva'] = 0.0
                        
                    synth['total'] = dr['imp_pagado']
                    synth['metodo_pago'] = f"REP (Origen: {parent['uuid'][:8]})"
                    flujo.append(synth)
                elif rep.get('total', 0.0) > 0:
                    flujo.append(rep)
        return flujo

    todas_emitidas = list(datos_emitidas)
    todas_recibidas = list(datos_recibidas)

    datos_emitidas = procesar_flujo(datos_emitidas)
    datos_recibidas = procesar_flujo(datos_recibidas)

    emit_event({"type": "progress", "percent": 65, "msg": "Calculando acumulados de IVA, ISR y cruces contra listas negras..."})
    time.sleep(0.5)

    emit_event({"type": "progress", "percent": 85, "msg": "Generando libro contable profesional en Excel (.xlsx)..."})
    generar_papel_trabajo_excel(
        datos_emitidas, datos_recibidas, args.year, args.zone, args.output,
        report_type=args.type, month=args.month, regime=args.regime,
        opcion_arrendamiento=args.arrendamiento_opcion
    )

    # Calcular métricas para el resumen ejecutivo
    d_em = [d for d in datos_emitidas if d.get('mes') == args.month] if (args.month and args.month > 0) else datos_emitidas
    d_rec = [d for d in datos_recibidas if d.get('mes') == args.month] if (args.month and args.month > 0) else datos_recibidas

    total_ing = sum(x['subtotal'] for x in d_em)
    # Deducciones autorizadas estrictas (excluyendo EFOS, Gasolina en efectivo y Efectivo > 2k)
    total_gas_deducible = sum(x.get('subtotal_deducible', x['subtotal']) for x in d_rec if x.get('es_deducible', True))
    total_gas_bruto = sum(x['subtotal'] for x in d_rec)
    total_gas_no_deducible = sum(x['subtotal'] for x in d_rec if not x.get('es_deducible', True))

    total_iva_t = sum(x['iva_16'] + x['iva_8'] for x in d_em)
    total_iva_a = sum(x['iva_16'] + x['iva_8'] for x in d_rec if x.get('es_deducible', True))
    total_ret_suf = sum(x['ret_iva'] for x in d_em)
    total_ret_isr = sum(x['ret_isr'] for x in d_em)
    saldo_iva = round(total_iva_t - total_iva_a - total_ret_suf, 2)

    regimen_global = calcular_impuestos_regimen(
        reg_code,
        total_ing,
        total_gas_deducible,
        retencion_isr=total_ret_isr,
        retencion_iva=total_ret_suf,
        opcion_arrendamiento=args.arrendamiento_opcion
    )
    resico_global = calcular_resico_mes(total_ing, ret_isr=total_ret_isr)

    if reg_code == '626':
        # En RESICO, los gastos no disminuyen la base del ISR (Art. 113-E LISR)
        utilidad_isr = regimen_global['isr_neto']
    else:
        utilidad_isr = regimen_global['base_isr']

    # Desglose mensual para gráficos y tablas en pantalla
    meses_a_emitir = [args.month] if (args.month and args.month > 0) else list(range(1, 13))
    mensual_data = []
    for m in meses_a_emitir:
        em_m = [x for x in datos_emitidas if x.get('mes') == m]
        rec_m = [x for x in datos_recibidas if x.get('mes') == m]
        s_em = sum(x['subtotal'] for x in em_m)
        v_16 = sum(x['subtotal'] for x in em_m if x['iva_16'] > 0)
        v_8 = sum(x['subtotal'] for x in em_m if x['iva_8'] > 0)
        iva_t = sum(x['iva_16'] + x['iva_8'] for x in em_m)

        s_rec = sum(x.get('subtotal_deducible', x['subtotal']) for x in rec_m if x.get('es_deducible', True))
        s_rec_no_ded = sum(x['subtotal'] for x in rec_m if not x.get('es_deducible', True))
        c_16 = sum(x.get('subtotal_deducible', x['subtotal']) for x in rec_m if x.get('es_deducible', True) and x['iva_16'] > 0)
        c_8 = sum(x.get('subtotal_deducible', x['subtotal']) for x in rec_m if x.get('es_deducible', True) and x['iva_8'] > 0)
        iva_a = sum(x['iva_16'] + x['iva_8'] for x in rec_m if x.get('es_deducible', True))
        r_iva = sum(x['ret_iva'] for x in em_m)
        r_isr = sum(x['ret_isr'] for x in em_m)
        s_iva = round(iva_t - iva_a - r_iva, 2)

        calc_reg_m = calcular_impuestos_regimen(
            reg_code, s_em, s_rec,
            retencion_isr=r_isr, retencion_iva=r_iva,
            opcion_arrendamiento=args.arrendamiento_opcion
        )
        res_m = calcular_resico_mes(s_em, ret_isr=r_isr)
        if reg_code == '626':
            u_isr = calc_reg_m['isr_neto']
        else:
            u_isr = calc_reg_m['base_isr']

        mensual_data.append({
            "mes": m,
            "nombre": MESES_NOMBRES[m],
            "ingresos": round(s_em, 2),
            "ventas_16": round(v_16, 2),
            "ventas_8": round(v_8, 2),
            "gastos": round(s_rec, 2),
            "gastos_no_deducibles": round(s_rec_no_ded, 2),
            "compras_16": round(c_16, 2),
            "compras_8": round(c_8, 2),
            "iva_trasladado": round(iva_t, 2),
            "iva_acreditable": round(iva_a, 2),
            "saldo_iva": s_iva,
            "ret_iva": round(r_iva, 2),
            "ret_isr": round(r_isr, 2),
            "utilidad_isr": round(u_isr, 2),
            "regimen_calc": calc_reg_m,
            "resico": res_m
        })

    # Concentrado de retenciones para la tabla en pantalla
    agrupado_ui = {}
    for d in d_em:
        if d['ret_isr'] > 0 or d['ret_iva'] > 0:
            key = ('Emitida (Cliente)', d['receptor_rfc'], d['receptor_nombre'])
            if key not in agrupado_ui:
                agrupado_ui[key] = {'subtotal': 0.0, 'ret_isr': 0.0, 'ret_iva': 0.0, 'count': 0}
            agrupado_ui[key]['subtotal'] += d['subtotal']
            agrupado_ui[key]['ret_isr'] += d['ret_isr']
            agrupado_ui[key]['ret_iva'] += d['ret_iva']
            agrupado_ui[key]['count'] += 1

    for d in d_rec:
        if d['ret_isr'] > 0 or d['ret_iva'] > 0:
            key = ('Recibida (Proveedor)', d['emisor_rfc'], d['emisor_nombre'])
            if key not in agrupado_ui:
                agrupado_ui[key] = {'subtotal': 0.0, 'ret_isr': 0.0, 'ret_iva': 0.0, 'count': 0}
            agrupado_ui[key]['subtotal'] += d['subtotal']
            agrupado_ui[key]['ret_isr'] += d['ret_isr']
            agrupado_ui[key]['ret_iva'] += d['ret_iva']
            agrupado_ui[key]['count'] += 1

    retenciones_data = []
    for (tipo_op, rfc_ter, nom_ter), vals in agrupado_ui.items():
        retenciones_data.append({
            "tipo": tipo_op,
            "rfc": rfc_ter,
            "nombre": nom_ter,
            "subtotal": round(vals['subtotal'], 2),
            "ret_isr": round(vals['ret_isr'], 2),
            "ret_iva": round(vals['ret_iva'], 2),
            "count": vals['count']
        })

    # Generar auditoría comprobante por comprobante de gastos y decisiones fiscales
    detalle_gastos = []
    efos_detectados = []
    seen_efos_rfc = set()

    d_rec_audit = [d for d in todas_recibidas if d.get('mes') == args.month] if (args.month and args.month > 0) else todas_recibidas

    for d in d_rec_audit:
        forma = str(d.get('forma_pago') or 'SIN_XML').strip()
        if forma and forma.isdigit():
            forma = forma.zfill(2)
        riesgo = str(d.get('riesgo_deducibilidad') or '')
        iva_ex = float(d.get('iva_exento') or 0.0)
        efecto = str(d.get('tipo') or 'I')
        is_nom = 'nomina' in riesgo.lower() or efecto.upper() == 'N' or 'nómina' in riesgo.lower() or d.get('es_nomina_emitida')
        no_ded_tipo = d.get('no_deducible_tipo') or ''
        is_gas = d.get('is_gasolina', False) or any(k in str(d.get('emisor_nombre', '')).lower() for k in ['gasolin', 'combust', 'gas popo', 'origas', 'gastoga', 'petro', 'diesel'])
        is_educ = any(w in riesgo.lower() for w in ['enseñanza', 'ensenan', 'educat', 'educac']) or any(w in str(d.get('emisor_nombre', '')).lower() for w in ['instituto', 'colegio', 'escuela', 'universidad', 'educat', 'educac', 'ensenan', 'liceo', 'kinder', 'primaria', 'secundaria', 'preparatoria', 'academ', 'formacion', 'docen'])
        is_med = any(w in riesgo.lower() for w in ['médic', 'medic', 'salud']) or any(w in str(d.get('emisor_nombre', '')).lower() for w in ['hospital', 'clinica', 'medico', 'sanatorio', 'salud', 'laboratorio', 'diagnostico'])
        is_rep = 'REP' in str(d.get('metodo_pago', '')).upper() or efecto.upper() == 'P' or efecto.upper() == 'PAGO' or no_ded_tipo == 'rep_control'
        is_egr = 'EGRESO' in str(d.get('metodo_pago', '')).upper() or efecto.upper() == 'E' or efecto.upper() == 'EGRESO'
        has_xml = bool(d.get('has_xml')) or (forma not in ['SIN_XML', 'sin_xml', 'SD', 'sd', ''] and forma != '99')

        if no_ded_tipo == 'efos_definitivo' or 'EFOS' in riesgo:
            decision = "No Deducible (EFOS Definitivo)"
            categoria = "efos"
            fundamento = "Art. 69-B CFF"
            explicacion = "El emisor figura en la lista definitiva del SAT de contribuyentes con operaciones inexistentes (EFOS). El comprobante carece de efectos fiscales (no deducible de ISR ni acreditable para IVA)."
            rfc_e = d.get('emisor_rfc', '')
            if rfc_e and rfc_e not in seen_efos_rfc:
                seen_efos_rfc.add(rfc_e)
                efos_detectados.append({
                    "rfc": rfc_e,
                    "nombre": d.get('emisor_nombre', ''),
                    "uuid": d.get('uuid', ''),
                    "fecha": d.get('fecha', ''),
                    "total": round(d.get('total', 0.0), 2),
                    "subtotal": round(d.get('subtotal', 0.0), 2),
                    "situacion": "Definitivo (Art. 69-B CFF)"
                })
        elif no_ded_tipo == 'gasolina_efectivo' or (is_gas and forma == '01'):
            decision = "No Deducible (Combustible en Efectivo)"
            categoria = "gasolina_efectivo"
            fundamento = "Art. 27 Fracc. III LISR"
            explicacion = "El combustible pagado en efectivo (01) es 100% no deducible sin importar el monto. Requiere tarjeta de débito/crédito, cheque nominativo, transferencia o monedero electrónico autorizado."
        elif no_ded_tipo == 'efectivo_mayor_2k' or (forma == '01' and d.get('total', 0.0) > 2000.0):
            decision = "No Deducible (Efectivo > $2,000 MXN)"
            categoria = "efectivo_2k"
            fundamento = "Art. 27 Fracc. III LISR"
            explicacion = "Los pagos que excedan $2,000.00 MXN liquidados en efectivo (01) no son deducibles ni permiten el acreditamiento de IVA. Requieren medios bancarizados."
        elif is_rep:
            decision = "Complemento para Recepción de Pagos (REP)"
            categoria = "rep"
            fundamento = "Art. 1-B LIVA & RMF 2.7.1.35"
            explicacion = "Ampara la liquidación bancarizada de una compra a crédito (PPD). Por disposición fiscal (Art. 1-B LIVA), su importe se mantiene en $0.00 para evitar duplicar la deducción del gasto que ya fue reconocido en la factura original."
        elif is_egr:
            decision = "Nota de Crédito de Proveedor (Disminuye Gasto)"
            categoria = "egreso"
            fundamento = "Art. 25 Fracc. VI LISR & Art. 7 LIVA"
            explicacion = "CFDI de egreso recibido de proveedor por devolución, descuento o bonificación sobre compras. Disminuye las deducciones autorizadas."
        elif is_nom:
            decision = "Sueldos y Salarios (Nómina - IVA $0.00)"
            categoria = "nomina"
            fundamento = "Art. 27 Fracc. V LISR & Arts. 14, 15 LIVA"
            explicacion = "CFDI de nómina deducible de personal para ISR; los servicios personales subordinados no causan ni trasladan IVA (IVA en $0.00)."
        elif is_educ or iva_ex > 0:
            decision = "Servicio de Enseñanza Exento de IVA"
            categoria = "exento"
            fundamento = "Art. 15 Fracc. IV LIVA"
            explicacion = "Servicio educativo con reconocimiento de validez oficial (RVOE). Deducible de ISR conforme al régimen; exento del traslado y acreditamiento de IVA."
        elif is_med:
            decision = "Honorarios Médicos Exentos de IVA"
            categoria = "exento"
            fundamento = "Art. 15 Fracc. XIV LIVA"
            explicacion = "Honorarios de servicios médicos o de salud exentos de IVA conforme al Art. 15 Fracc. XIV de la LIVA."
        elif is_gas and not has_xml:
            decision = "Condicionado a Pago Electrónico (Combustible)"
            categoria = "condicionado_bancarizacion"
            fundamento = "Art. 27 Fracc. III LISR"
            explicacion = "El combustible requiere forzosamente pago con tarjeta, transferencia, cheque o monedero electrónico autorizado. Registro procedente de metadatos; valida con el XML timbrado o estado de cuenta que no se haya pagado en efectivo (01)."
        elif d.get('total', 0.0) > 2000.0 and not has_xml:
            decision = "Condicionado a Pago Bancarizado (> $2,000)"
            categoria = "condicionado_bancarizacion"
            fundamento = "Art. 27 Fracc. III LISR"
            explicacion = "Los pagos que excedan $2,000.00 MXN requieren medios bancarizados (tarjeta, transferencia o cheque nominativo). Factura en PUE (Contado); corrobora con tu estado de cuenta bancario que no haya sido liquidada en efectivo (01)."
        elif d.get('metodo_pago') == 'PPD':
            decision = "Pendiente de Complemento (PPD)"
            categoria = "diferido"
            fundamento = "Art. 1-B LIVA"
            explicacion = "Factura emitida a crédito o diferida; el acreditamiento requiere la expedición del Complemento de Pago (REP)."
        else:
            decision = "100% Deducible e IVA Acreditable"
            categoria = "deducible"
            fundamento = "Art. 27 LISR & Art. 5 LIVA"
            explicacion = "Cumple con requisitos fiscales indispensables, pago bancarizado o permitido en efectivo (< $2,000) y proveedor vigente no listado en EFOS."

        m_idx = d.get('mes', 1)
        m_name = MESES_NOMBRES[m_idx] if (0 < m_idx < 13) else f"Mes {m_idx}"
        detalle_gastos.append({
            "uuid": d.get('uuid', ''),
            "fecha": d.get('fecha', ''),
            "mes": m_idx,
            "mes_nombre": m_name,
            "rfc": d.get('emisor_rfc', ''),
            "nombre": d.get('emisor_nombre', ''),
            "metodo_pago": d.get('metodo_pago', 'PUE'),
            "forma_pago": forma,
            "forma_pago_nombre": get_forma_pago_nombre(forma, d.get('metodo_pago', 'PUE')),
            "subtotal": round(d.get('subtotal', 0.0), 2),
            "subtotal_deducible": round(d.get('subtotal_deducible', d.get('subtotal', 0.0)), 2),
            "es_deducible": d.get('es_deducible', True),
            "iva_16": round(d.get('iva_16', 0.0), 2),
            "iva_8": round(d.get('iva_8', 0.0), 2),
            "iva_exento": round(iva_ex, 2),
            "total": round(d.get('total', 0.0), 2),
            "decision": decision,
            "categoria": categoria,
            "fundamento": fundamento,
            "explicacion": explicacion,
            "riesgo": riesgo
        })

    file_size = os.path.getsize(args.output) if os.path.exists(args.output) else 0
    event_data = {
        "type": "completed",
        "percent": 100,
        "msg": f"Reporte de Análisis de IVA e ISR ({reg_info.get('short_label', reg_code)}) generado exitosamente.",
        "path": os.path.abspath(args.output),
        "filename": os.path.basename(args.output),
        "size": file_size,
        "metrics": {
            "has_real_data": has_real_data,
            "total_emitidas": len(todas_emitidas),
            "total_recibidas": len(todas_recibidas),
            "ingresos": round(total_ing, 2),
            "gastos": round(total_gas_deducible, 2),
            "gastos_brutos": round(total_gas_bruto, 2),
            "gastos_no_deducibles": round(total_gas_no_deducible, 2),
            "iva_trasladado": round(total_iva_t, 2),
            "iva_acreditable": round(total_iva_a, 2),
            "saldo_iva": saldo_iva,
            "utilidad_isr": utilidad_isr,
            "retencion_iva": round(total_ret_suf, 2),
            "retencion_isr": round(total_ret_isr, 2),
            "regime": reg_code,
            "regime_code": reg_code,
            "regime_info": reg_info,
            "regimen_calc": regimen_global,
            "is_resico": reg_code == '626',
            "resico": {
                "base_anual": round(total_ing, 2),
                "tasa": resico_global['tasa'],
                "tasa_porcentaje": resico_global['tasa_porcentaje'],
                "isr_causado": resico_global['isr_causado'],
                "retencion_isr": resico_global['retencion_125'],
                "isr_neto": resico_global['isr_neto'],
                "isr_a_favor": resico_global['isr_a_favor']
            },
            "efos_detectados": efos_detectados,
            "mensual": mensual_data,
            "retenciones": retenciones_data,
            "detalle_gastos": detalle_gastos
        }
    }

    # Guardar métricas directamente a un archivo JSON complementario para evitar saturación de buffer de stdout
    try:
        metrics_file = args.metrics_json or f"{args.output}.metrics.json"
        with open(metrics_file, "w", encoding="utf-8") as mf:
            json.dump(event_data, mf, ensure_ascii=False)
    except Exception:
        pass

    emit_event(event_data)

if __name__ == "__main__":
    main()
