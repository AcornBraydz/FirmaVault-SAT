import os
import sys
import argparse
import json
import zipfile
from defusedxml import ElementTree as ET
from datetime import datetime
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def emit_event(data):
    try:
        print(json.dumps(data), flush=True)
    except Exception:
        pass

def parse_nomina_xml(xml_content):
    try:
        root = ET.fromstring(xml_content)
    except Exception:
        return None

    tipo_comp = root.attrib.get('TipoDeComprobante', '')
    if tipo_comp != 'N' and tipo_comp != 'E':
        has_nomina = any('Nomina' in elem.tag for elem in root.iter())
        if not has_nomina and tipo_comp != 'N':
            return None
    
    fecha_raw = root.attrib.get('Fecha', '')
    folio_fiscal = ""
    for elem in root.iter():
        tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
        if tag == 'TimbreFiscalDigital':
            folio_fiscal = elem.attrib.get('UUID', '')

    receptor_rfc, receptor_nombre = "", ""
    for elem in root:
        tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
        if tag == 'Receptor':
            receptor_rfc = elem.attrib.get('Rfc', '')
            receptor_nombre = elem.attrib.get('Nombre', receptor_rfc)

    # Find Nomina node
    nomina_node = None
    for elem in root.iter():
        tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
        if tag == 'Nomina':
            nomina_node = elem
            break

    if nomina_node is None:
        return None

    fecha_pago = nomina_node.attrib.get('FechaPago', fecha_raw)
    dt = None
    if fecha_pago:
        try:
            dt = datetime.fromisoformat(fecha_pago.split("T")[0])
        except Exception:
            pass

    gravado = 0.0
    exento = 0.0
    isr = 0.0
    subsidio_empleo = 0.0
    tipo_regimen = '02'
    tipo_nomina = nomina_node.attrib.get('TipoNomina', 'O')

    for elem in root.iter():
        tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
        if tag == 'Receptor' and 'TipoRegimen' in elem.attrib:
            tipo_regimen = elem.attrib.get('TipoRegimen', '02')
        elif tag == 'Percepcion':
            grav = float(elem.attrib.get('ImporteGravado', 0.0))
            exe = float(elem.attrib.get('ImporteExento', 0.0))
            gravado += grav
            exento += exe
        elif tag == 'Deduccion':
            tipo_deduccion = elem.attrib.get('TipoDeduccion', '')
            if tipo_deduccion in ['002', '2']:
                isr += float(elem.attrib.get('Importe', 0.0))
        elif tag == 'OtroPago':
            top = elem.attrib.get('TipoOtroPago', '')
            if top in ['002', '2']: # Subsidio al Empleo (SPE) entregado
                imp_spe = float(elem.attrib.get('Importe', 0.0))
                subsidio_empleo += imp_spe
                for ch in elem.iter():
                    if 'SubsidioCausado' in ch.attrib:
                        try:
                            subsidio_empleo = max(subsidio_empleo, float(ch.attrib.get('SubsidioCausado', 0.0)))
                        except Exception:
                            pass

    return {
        'uuid': folio_fiscal,
        'fecha_pago': fecha_pago,
        'mes': dt.month if dt else 1,
        'anio': dt.year if dt else 2026,
        'receptor_rfc': receptor_rfc,
        'receptor_nombre': receptor_nombre,
        'gravado': round(gravado, 2),
        'exento': round(exento, 2),
        'isr': round(isr, 2),
        'subsidio_empleo': round(subsidio_empleo, 2),
        'isr_neto': round(max(0.0, isr - subsidio_empleo), 2),
        'tipo_regimen': tipo_regimen,
        'tipo_nomina': tipo_nomina,
    }

def build_excel(datos, output_path, anio, mes):
    wb = Workbook()
    
    color_header = "0F766E" # Teal 700
    color_total = "CCFBF1" # Teal 100
    
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_regular = Font(name="Calibri", size=10)
    
    fill_header = PatternFill(start_color=color_header, end_color=color_header, fill_type="solid")
    fill_total = PatternFill(start_color=color_total, end_color=color_total, fill_type="solid")
    
    thin_border = Border(left=Side(style='thin', color='E2E8F0'), 
                         right=Side(style='thin', color='E2E8F0'), 
                         top=Side(style='thin', color='E2E8F0'), 
                         bottom=Side(style='thin', color='E2E8F0'))
                         
    num_format = '"$"#,##0.00_-'

    # Aggregating data
    empleados = defaultdict(lambda: {'nombre': '', 'gravado': 0.0, 'exento': 0.0, 'isr': 0.0, 'recibos': 0})
    total_gravado = 0.0
    total_exento = 0.0
    total_isr = 0.0

    total_subsidio = 0.0

    for d in datos:
        rfc = d['receptor_rfc']
        empleados[rfc]['nombre'] = d['receptor_nombre']
        empleados[rfc]['gravado'] += d['gravado']
        empleados[rfc]['exento'] += d['exento']
        empleados[rfc]['isr'] += d['isr']
        empleados[rfc]['subsidio'] = empleados[rfc].get('subsidio', 0.0) + d.get('subsidio_empleo', 0.0)
        empleados[rfc]['isr_neto'] = max(0.0, empleados[rfc]['isr'] - empleados[rfc]['subsidio'])
        empleados[rfc]['recibos'] += 1
        
        total_gravado += d['gravado']
        total_exento += d['exento']
        total_isr += d['isr']
        total_subsidio += d.get('subsidio_empleo', 0.0)

    total_neto = max(0.0, total_isr - total_subsidio)

    # ---------------------------------------------------------
    # Tab 1: Concentrado Empresa
    # ---------------------------------------------------------
    ws_resumen = wb.active
    ws_resumen.title = "1_Concentrado_Declaracion"
    ws_resumen.views.sheetView[0].showGridLines = False
    
    ws_resumen.append(["Resumen de Nóminas para Declarar (Cédula de Retenciones)"])
    ws_resumen.cell(row=1, column=1).font = Font(size=14, bold=True, color=color_header)
    
    periodo = f"Mes: {mes} | Año: {anio}" if mes > 0 else f"Anual Consolidado | Año: {anio}"
    ws_resumen.append([periodo])
    ws_resumen.append([])
    
    headers_res = ["Concepto SAT (Declaraciones y Pagos)", "Importe Total"]
    ws_resumen.append(headers_res)
    for c in range(1, 3):
        ws_resumen.cell(row=4, column=c).font = font_header
        ws_resumen.cell(row=4, column=c).fill = fill_header
        ws_resumen.cell(row=4, column=c).alignment = Alignment(horizontal="center")

    rows = [
        ["Total Ingresos Gravados por Sueldos y Salarios", total_gravado],
        ["Total Ingresos Exentos por Sueldos y Salarios", total_exento],
        ["Total Percepciones Brutas (Gravado + Exento)", total_gravado + total_exento],
        ["Total ISR Retenido por Salarios (Deducción 002)", total_isr],
        ["Subsidio para el Empleo Causado/Entregado (SPE)", total_subsidio],
        ["ISR Neto a Cargo para Pago Provisional", total_neto]
    ]
    
    row_idx = 5
    for r in rows:
        ws_resumen.append(r)
        c1 = ws_resumen.cell(row=row_idx, column=1)
        c2 = ws_resumen.cell(row=row_idx, column=2)
        c1.font = font_bold
        c2.number_format = num_format
        c2.font = font_bold
        c1.border = thin_border
        c2.border = thin_border
        if "ISR Neto a Cargo" in r[0]:
            c1.fill = fill_total
            c2.fill = fill_total
        row_idx += 1

    ws_resumen.column_dimensions['A'].width = 46
    ws_resumen.column_dimensions['B'].width = 25

    # ---------------------------------------------------------
    # Tab 2: Desglose por Empleado
    # ---------------------------------------------------------
    ws_emp = wb.create_sheet(title="2_Desglose_Trabajadores")
    ws_emp.views.sheetView[0].showGridLines = True
    
    headers_emp = ["RFC Trabajador", "Nombre del Trabajador", "No. Recibos", "Ingreso Gravado", "Ingreso Exento", "ISR Retenido (002)", "Subsidio SPE", "ISR Neto"]
    ws_emp.append(headers_emp)
    for col_idx in range(1, len(headers_emp) + 1):
        cell = ws_emp.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 2
    for rfc, info in sorted(empleados.items(), key=lambda x: x[1]['nombre']):
        ws_emp.append([
            rfc,
            info['nombre'],
            info['recibos'],
            info['gravado'],
            info['exento'],
            info['isr'],
            info.get('subsidio', 0.0),
            info.get('isr_neto', 0.0)
        ])
        for c in range(1, 9):
            cell = ws_emp.cell(row=row_idx, column=c)
            cell.border = thin_border
            cell.font = font_regular
            if c >= 4:
                cell.number_format = num_format
            if c == 3:
                cell.alignment = Alignment(horizontal="center")
        row_idx += 1

    # Add total row at the bottom
    ws_emp.append(["TOTALES", "", len(datos), total_gravado, total_exento, total_isr, total_subsidio, total_neto])
    for c in range(1, 9):
        cell = ws_emp.cell(row=row_idx, column=c)
        cell.font = font_bold
        cell.fill = fill_total
        if c >= 4:
            cell.number_format = num_format
            
    for col in ws_emp.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if not val.startswith("="):
                max_len = max(max_len, len(val))
        ws_emp.column_dimensions[col_letter].width = max(max_len + 3, 12)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--year', type=int, required=True)
    parser.add_argument('--month', type=int, required=False, default=0)
    parser.add_argument('--xml-dir', type=str, required=False, default='')
    parser.add_argument('--rfc', type=str, required=False, default='')
    parser.add_argument('--output', type=str, required=False, default='')
    parser.add_argument('--preview', action='store_true', default=False)
    parser.add_argument('--json-output', type=str, required=False, default='')
    args = parser.parse_args()

    emit_event({"type": "progress", "percent": 10, "msg": "Iniciando análisis de nóminas..."})

    datos = []
    search_dirs = []
    
    if args.xml_dir and os.path.exists(args.xml_dir):
        search_dirs.append(args.xml_dir)

    base_tools_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
    
    # 1. Scratch xml_nominas
    scratch_dir = os.path.join(base_tools_dir, "scratch", "xml_nominas")
    if os.path.exists(scratch_dir):
        search_dirs.append(scratch_dir)

    # 2. Tools nominas si viene RFC
    if args.rfc:
        client_nom_dir = os.path.join(base_tools_dir, "tools", "nominas", args.rfc.upper().strip())
        if os.path.exists(client_nom_dir):
            search_dirs.append(client_nom_dir)

    # 3. Scratch xml_storage
    scratch_storage = os.path.join(base_tools_dir, "scratch", "xml_storage")
    if os.path.exists(scratch_storage):
        search_dirs.append(scratch_storage)

    processed = 0
    seen_uuids = set()

    def process_content(content):
        nonlocal processed
        parsed = parse_nomina_xml(content)
        if parsed and parsed.get('uuid'):
            u = parsed['uuid'].upper().strip()
            if u in seen_uuids:
                return
            seen_uuids.add(u)

            if args.month != 0 and parsed['mes'] != args.month:
                return
            if parsed['anio'] != args.year:
                return
            datos.append(parsed)
            
            processed += 1
            if processed % 50 == 0:
                emit_event({"type": "progress", "percent": 15 + int((processed / 100) * 50), "msg": f"Procesando XMLs de Nómina ({processed})..."})

    for target_dir in search_dirs:
        for root_dir, dirs, files in os.walk(target_dir):
            for file in files:
                file_path = os.path.join(root_dir, file)
                if file.lower().endswith(".xml"):
                    try:
                        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                            content = f.read()
                        process_content(content)
                    except Exception:
                        pass
                elif file.lower().endswith(".zip"):
                    try:
                        with zipfile.ZipFile(file_path, 'r') as z:
                            for zname in z.namelist():
                                if zname.lower().endswith(".xml"):
                                    try:
                                        with z.open(zname) as zf:
                                            content = zf.read().decode('utf-8', errors='ignore')
                                        process_content(content)
                                    except Exception:
                                        pass
                    except Exception:
                        pass
                        
    if processed == 0:
        emit_event({"type": "progress", "percent": 50, "msg": "No se encontraron recibos de nómina locales."})

    emit_event({"type": "progress", "percent": 75, "msg": "Consolidando cifras para declaración SAT..."})
    
    empleados_map = defaultdict(lambda: {'nombre': '', 'gravado': 0.0, 'exento': 0.0, 'isr': 0.0, 'subsidio': 0.0, 'isr_neto': 0.0, 'recibos': 0})
    mensual_map = defaultdict(lambda: {'gravado': 0.0, 'exento': 0.0, 'isr': 0.0, 'subsidio': 0.0, 'isr_neto': 0.0, 'recibos': 0})

    total_gravado = 0.0
    total_exento = 0.0
    total_isr = 0.0
    total_subsidio = 0.0
    total_asimilados = 0.0
    total_sueldos = 0.0

    for d in datos:
        rfc = d['receptor_rfc']
        empleados_map[rfc]['nombre'] = d['receptor_nombre']
        empleados_map[rfc]['gravado'] += d['gravado']
        empleados_map[rfc]['exento'] += d['exento']
        empleados_map[rfc]['isr'] += d['isr']
        empleados_map[rfc]['subsidio'] += d.get('subsidio_empleo', 0.0)
        empleados_map[rfc]['recibos'] += 1

        m_num = d.get('mes', 1)
        mensual_map[m_num]['gravado'] += d['gravado']
        mensual_map[m_num]['exento'] += d['exento']
        mensual_map[m_num]['isr'] += d['isr']
        mensual_map[m_num]['subsidio'] += d.get('subsidio_empleo', 0.0)
        mensual_map[m_num]['recibos'] += 1

        total_gravado += d['gravado']
        total_exento += d['exento']
        total_isr += d['isr']
        total_subsidio += d.get('subsidio_empleo', 0.0)

        if str(d.get('tipo_regimen', '')).startswith('05') or str(d.get('tipo_regimen', '')).startswith('06') or str(d.get('tipo_regimen', '')).startswith('07'):
            total_asimilados += d['gravado'] + d['exento']
        else:
            total_sueldos += d['gravado'] + d['exento']

    for rfc, v in empleados_map.items():
        v['isr_neto'] = round(max(0.0, v['isr'] - v['subsidio']), 2)

    MESES_NOMBRES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
    mensual_arr = []
    for m_i in range(1, 13):
        m_data = mensual_map[m_i]
        mensual_arr.append({
            "mes": m_i,
            "nombre": MESES_NOMBRES[m_i],
            "gravado": round(m_data['gravado'], 2),
            "exento": round(m_data['exento'], 2),
            "percepciones": round(m_data['gravado'] + m_data['exento'], 2),
            "isr": round(m_data['isr'], 2),
            "subsidio": round(m_data['subsidio'], 2),
            "isr_neto": round(max(0.0, m_data['isr'] - m_data['subsidio']), 2),
            "imss": 0.0,
            "recibos": m_data['recibos']
        })

    trabajadores_arr = []
    for rfc, v in sorted(empleados_map.items(), key=lambda x: x[1]['nombre']):
        trabajadores_arr.append({
            "rfc": rfc,
            "nombre": v['nombre'],
            "gravado": round(v['gravado'], 2),
            "exento": round(v['exento'], 2),
            "percepciones": round(v['gravado'] + v['exento'], 2),
            "isr": round(v['isr'], 2),
            "subsidio": round(v['subsidio'], 2),
            "isr_neto": v['isr_neto'],
            "imss": 0.0,
            "recibos": v['recibos']
        })

    preview_payload = {
        "metricas": {
            "total_trabajadores": len(empleados_map),
            "total_recibos": len(datos),
            "total_gravado": round(total_gravado, 2),
            "total_exento": round(total_exento, 2),
            "total_percepciones": round(total_gravado + total_exento, 2),
            "total_isr_retenido": round(total_isr, 2),
            "total_subsidio_empleo": round(total_subsidio, 2),
            "isr_neto_a_pagar": round(max(0.0, total_isr - total_subsidio), 2),
            "sueldos_ordinarios": round(total_sueldos, 2),
            "asimilados": round(total_asimilados, 2)
        },
        "mensual": mensual_arr,
        "trabajadores": trabajadores_arr
    }

    if args.json_output:
        try:
            os.makedirs(os.path.dirname(os.path.abspath(args.json_output)), exist_ok=True)
            with open(args.json_output, 'w', encoding='utf-8') as f:
                json.dump(preview_payload, f, ensure_ascii=False, indent=2)
        except Exception as e:
            pass

    if args.output:
        build_excel(datos, args.output, args.year, args.month)

    emit_event({"type": "progress", "percent": 100, "msg": "Generación completa"})
    if args.output:
        emit_event({"type": "finished", "excel_path": args.output, "preview": preview_payload})
    else:
        emit_event({"type": "preview_result", "data": preview_payload})

if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        emit_event({"type": "error", "error": str(e)})
        sys.exit(1)
